"""Synthetic fixtures only: review lifecycle, no-ZIP freezes and native Word tables."""
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Border, Side

from wfcore import registry, reviewrounds as rr, scientificfreeze as sf
from wfcore.readiness import atomic_json, verify_reader_review
from wfcore.state import State
from wfcore.wordtables import append_tables, table_problems, load_spec
from manuscript.build_docx import normalize_docx, audit_docx
from test_source_aligned import DEFAULT_STYLE


class ReviewRoundTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="medpaper-review-rounds-")
        self.addCleanup(self.temp.cleanup)
        self.project = Path(self.temp.name)
        self.state = State(self.project).create("medpaper", "1.12.0", "S19_human_review")
        self.data("01_protocol/artifact_plan.json", {"main_figures": [], "main_tables": []})
        for rel in ("07_manuscript/full_manuscript.md", "05_figures/legends.md", "04_tables/table_captions.md"):
            self.text(rel, "Synthetic fixture only, not a real research manuscript.")

    def text(self, rel, text):
        path = self.project / rel; path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8"); return path

    def data(self, rel, value):
        atomic_json(self.project / rel, value); return self.project / rel

    def report(self, record, *, current=False):
        selection = rr.inputs(self.project, record["scope"]) if current else record["reviewed_artifacts"]
        headings = ("Reader comprehension", "Medical presentation") if record["scope"] == "manuscript" else (
            "Reader comprehension", "Editorial screening and low-level errors", "Guideline compliance")
        text = "# Synthetic reviewer fixture\n\n" + "\n".join(
            "## " + heading + "\n\n" + ("Test-only review findings for controlled synthetic content, not actual independent scientific review. " * 2)
            for heading in headings)
        text += "\n```json\n" + json.dumps({"reviewed_artifacts": selection,"findings":[{"id":"C1","category":"preference","affects_science":False}]}) + "\n```\n"
        return self.text("temp/actual-test-review.md", text)

    def ready(self, scope="manuscript"):
        record = rr.start(self.project, scope)
        rr.attach(self.project, scope, self.report(record))
        return rr.close(self.project, scope, [], [])

    def test_close_without_report_is_rejected(self):
        rr.start(self.project, "manuscript")
        with self.assertRaises(ValueError): rr.close(self.project, "manuscript", [], [])

    def test_drift_while_reading_and_explicit_recovery(self):
        record = rr.start(self.project, "manuscript"); report = self.report(record)
        self.text("07_manuscript/full_manuscript.md", "Changed synthetic paper")
        with self.assertRaises(ValueError): rr.attach(self.project, "manuscript", report)
        restarted = rr.restart(self.project, "manuscript", "Actual input drift while synthetic reviewer was reading")
        self.assertEqual(restarted["id"], "M002")
        self.assertEqual(json.loads((self.project/rr.ROOT/"M001.json").read_text())["user_decision"]["value"], "abandoned")

    def test_round_freeze_has_no_zip_and_survives_package_round(self):
        self.ready(); rr.decide(self.project, "manuscript", "no_further_review", "我确认这一版")
        _, payload, _ = sf.write(self.project, "1.12.0")
        self.assertEqual(payload["schema_version"], 5)
        self.assertEqual(payload["review_round_id"], "M001")
        self.assertFalse(any("review_packages" in row["path"] for row in payload["files"]))
        self.assertEqual(list(self.project.rglob("*.zip")), [])
        self.assertTrue(sf.verify(self.project)[0])
        self.package(); rr.start(self.project, "package")
        self.assertTrue(sf.verify(self.project)[0])

    def test_post_closure_drift_blocks_user_answer(self):
        self.ready(); self.text("05_figures/legends.md", "Changed legend")
        with self.assertRaises(ValueError): rr.decide(self.project, "manuscript", "no_further_review", "我确认这一版")

    def test_generic_continue_and_missing_history_cannot_approve(self):
        self.ready()
        with self.assertRaises(ValueError): rr.decide(self.project, "manuscript", "no_further_review", "继续")
        rr.decide(self.project, "manuscript", "no_further_review", "我确认这一版")
        record = rr.latest(self.project, "manuscript")
        self.text(record["report"]["path"], "Tampered synthetic review")
        with self.assertRaises(ValueError): rr.approved(self.project, "manuscript")

    def test_unresolved_defect_blocks_approval_but_allows_another_round(self):
        record = rr.start(self.project, "manuscript")
        rr.attach(self.project, "manuscript", self.report(record))
        rr.close(self.project, "manuscript", [], [["C1", "Missing actual author-owned fact"]], 1)
        with self.assertRaises(ValueError): rr.decide(self.project, "manuscript", "no_further_review", "我确认这一版")
        rr.decide(self.project, "manuscript", "another_round", "再审一轮")
        self.assertEqual(rr.start(self.project, "manuscript")["id"], "M002")

    def test_changed_artifacts_must_be_recorded_and_revalidation_is_real(self):
        record = rr.start(self.project, "manuscript")
        rr.attach(self.project, "manuscript", self.report(record))
        self.text("07_manuscript/full_manuscript.md", "Changed synthetic scientific statement")
        with self.assertRaises(ValueError): rr.close(self.project, "manuscript", [], [])
        self.assertTrue(verify_reader_review(self.project))
        rr.attach(self.project, "manuscript", self.report(record, current=True), revalidation=True)
        self.assertEqual(verify_reader_review(self.project), [])
        rr.close(self.project, "manuscript", [["07_manuscript/full_manuscript.md", "Actual changed test statement validated"]], [])
        self.assertEqual(len(rr.latest(self.project, "manuscript")["reports"]), 2)

    def test_delta_detects_deletion_and_dependency_failure_widens(self):
        self.text("07_manuscript/supplementary_methods.md", "Optional synthetic methods")
        self.ready(); rr.decide(self.project, "manuscript", "another_round", "再审一轮")
        (self.project/"07_manuscript/supplementary_methods.md").unlink()
        with patch("wfcore.dependencies.closure", side_effect=ValueError("malformed dependency")):
            record = rr.start(self.project, "manuscript", delta=True)
        self.assertIn("07_manuscript/supplementary_methods.md", record["deleted_since_previous"])
        self.assertEqual(record["reviewed_artifacts"], record["before"])

    def test_new_round_invalidates_old_approval(self):
        self.ready(); rr.decide(self.project, "manuscript", "no_further_review", "我确认这一版")
        rr.start(self.project, "manuscript")
        self.assertIsNone(State(self.project).load().decision("manuscript_human_reviewed"))
        with self.assertRaises(ValueError): rr.approved(self.project, "manuscript")

    def package(self):
        self.text("08_submission/bundle/test.txt", "Synthetic upload")
        self.data("08_submission/bundle/manifest.json", {"items": [{"role": "test", "file": "08_submission/bundle/test.txt"}]})
        self.state.data["current"] = "S24_package_human_review"; self.state.save()

    def test_package_round_binds_requirements(self):
        self.package(); self.text("08_submission/guidelines_extract.md", "Current synthetic rule")
        self.ready("package")
        self.text("08_submission/guidelines_extract.md", "Changed synthetic rule")
        with self.assertRaises(ValueError): rr.decide(self.project, "package", "no_further_review", "我确认这一版")

    def test_empty_manifest_is_not_a_reviewable_package(self):
        self.data("08_submission/bundle/manifest.json", {"items": []})
        with self.assertRaises(ValueError): rr.start(self.project, "package")


class NativeWordTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="medpaper-native-word-"); self.addCleanup(self.temp.cleanup)
        self.project = Path(self.temp.name)

    def workbook(self, columns=3, merged=False):
        path = self.project / "table.xlsx"
        workbook = Workbook(); workbook.active.title = "Unrelated"
        workbook.active.append(["Never use the first sheet silently"])
        sheet = workbook.create_sheet("Actual table")
        sheet.append(["Table 1. Synthetic table"])
        sheet.append([f"Column {i+1}" for i in range(columns)])
        if merged:
            sheet.merge_cells(start_row=2, start_column=2, end_row=2, end_column=3)
            sheet.append(["Measure"] + [f"Group {i}" for i in range(1, columns)])
        sheet.append(["Synthetic measure"] + [i for i in range(1, columns)])
        sheet.append(["Second measure"] + [0 for _ in range(1, columns)])
        last = sheet.max_row
        for cell in sheet[last]: cell.border = Border(bottom=Side(style="thin"))
        sheet.append(["Values are synthetic test fixtures, not clinical findings."])
        workbook.save(path); workbook.close()
        return {"id": "Table 1", "file": "table.xlsx", "sheet": "Actual table", "header_start": 2,
                "header_end": 3 if merged else 2, "body_end": last, "orientation": "auto"}

    def test_native_three_lines_and_notes_outside_grid(self):
        entry = self.workbook(); doc = Document(); doc.add_paragraph("Synthetic QA specimen")
        append_tables(doc, self.project, [entry], DEFAULT_STYLE)
        self.assertEqual(table_problems(doc), [])
        self.assertEqual(len(doc.tables[0].rows), 3)
        self.assertNotIn("Never use", " ".join(p.text for p in doc.paragraphs))
        self.assertIn("Values are synthetic", doc.paragraphs[-1].text)
        self.assertEqual(doc.tables[0].cell(2, 1).text, "0")

    def test_wide_then_portrait_and_finalize_preserves_geometry(self):
        wide = self.workbook(columns=9); narrow = dict(wide, orientation="portrait")
        doc = Document(); doc.add_paragraph("Synthetic mixed orientation specimen")
        append_tables(doc, self.project, [wide, narrow], DEFAULT_STYLE)
        path = self.project/"tables.docx"; doc.save(path)
        normalize_docx(path, DEFAULT_STYLE, "tables"); normalize_docx(path, DEFAULT_STYLE, "tables")
        final = Document(path)
        self.assertEqual([s.orientation for s in final.sections], [WD_ORIENT.PORTRAIT, WD_ORIENT.LANDSCAPE, WD_ORIENT.PORTRAIT])
        self.assertEqual(table_problems(final), [])
        self.assertEqual(len(final.sections[0].footer._element.findall(".//"+qn("w:fldSimple"))), 1)
        self.assertEqual(audit_docx(path, DEFAULT_STYLE, "tables"), [])

    def test_merged_header_is_preserved(self):
        entry = self.workbook(merged=True); doc = Document(); doc.add_paragraph("Synthetic specimen")
        append_tables(doc, self.project, [entry], DEFAULT_STYLE)
        self.assertTrue(doc.tables[0]._tbl.xpath(".//w:gridSpan"))
        self.assertEqual(table_problems(doc), [])

    def test_ambiguous_footnotes_and_formulas_fail_closed(self):
        entry = self.workbook(); workbook = __import__("openpyxl").load_workbook(self.project/"table.xlsx")
        sheet = workbook[entry["sheet"]]; sheet.cell(sheet.max_row, 2, "Not a note")
        workbook.save(self.project/"table.xlsx"); workbook.close()
        with self.assertRaises(ValueError): append_tables(Document(), self.project, [entry], DEFAULT_STYLE)
        entry = self.workbook(); workbook = __import__("openpyxl").load_workbook(self.project/"table.xlsx")
        workbook[entry["sheet"]].cell(3, 2, "=1+1"); workbook.save(self.project/"table.xlsx"); workbook.close()
        with self.assertRaises(ValueError): append_tables(Document(), self.project, [entry], DEFAULT_STYLE)

    def test_sheet_is_explicit_and_copies_include_all_sheets(self):
        entry = self.workbook(); spec = self.project/"spec.json"
        atomic_json(spec, {"tables": [entry]}); self.assertEqual(len(load_spec(self.project, spec)[1]), 2)
        entry.pop("sheet"); atomic_json(spec, {"tables": [entry]})
        with self.assertRaises(ValueError): load_spec(self.project, spec)
        from wfcore.packagefreeze import _sha256
        copies = rr.reader_copies(self.project, {"table.xlsx": _sha256(self.project/"table.xlsx")}, ".wf/test-copies")
        text = (self.project/copies["table.xlsx"]["path"]).read_text()
        self.assertIn("Sheet: Unrelated", text); self.assertIn("Sheet: Actual table", text)

    def test_excel_percent_format_cannot_silently_change_units(self):
        entry = self.workbook(); workbook = __import__("openpyxl").load_workbook(self.project/"table.xlsx")
        cell = workbook[entry["sheet"]].cell(3, 2, 0.12); cell.number_format = "0.0%"
        workbook.save(self.project/"table.xlsx"); workbook.close()
        with self.assertRaisesRegex(ValueError, "units or precision"):
            append_tables(Document(), self.project, [entry], DEFAULT_STYLE)

    def test_reference_spacing_is_not_tab_unless_explicit(self):
        path = self.project/"reference.docx"; doc = Document(); doc.add_paragraph("Synthetic reference fixture")
        from docx.enum.style import WD_STYLE_TYPE
        doc.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
        doc.add_paragraph("1.\tSynthetic author. Test-only reference.", style="Bibliography")
        doc.save(path); normalize_docx(path, DEFAULT_STYLE, "tables")
        self.assertTrue(Document(path).paragraphs[1].text.startswith("1. Synthetic"))
        tab_style = {**DEFAULT_STYLE, "reference_separator": "tab"}
        normalize_docx(path, tab_style, "tables")
        self.assertTrue(Document(path).paragraphs[1].text.startswith("1.\tSynthetic"))

    def test_measurements_bind_actual_final_files_and_reuse_without_reaudit(self):
        from wfcore import submissionqc as qc
        output = self.project/"08_submission/bundle/tables.docx"; output.parent.mkdir(parents=True)
        doc = Document(); doc.add_paragraph("Synthetic measured fixture"); doc.save(output)
        normalize_docx(output, DEFAULT_STYLE, "tables")
        atomic_json(self.project/"08_submission/bundle/manifest.json", {"items": [
            {"role": "tables", "file": "08_submission/bundle/tables.docx"}]})
        atomic_json(self.project/"08_submission/docx_style.json", DEFAULT_STYLE)
        checklist = self.project/qc.CHECKLIST; checklist.parent.mkdir(parents=True)
        checklist.write_text("# Manual findings\n\nActual visual inspection remains separate.\n", encoding="utf-8")
        with patch("wfcore.submission.verify_portal", return_value=[]):
            self.assertTrue(qc.write(self.project)["ok"])
        with patch.object(qc, "measure", side_effect=AssertionError("unchanged files must reuse measured evidence")):
            self.assertEqual(qc.verify(self.project), [])
        self.assertIn("Actual visual inspection remains separate.", checklist.read_text())
        output.write_bytes(output.read_bytes()+b"drift")
        self.assertTrue(qc.verify(self.project))


if __name__ == "__main__":
    unittest.main(verbosity=2)
