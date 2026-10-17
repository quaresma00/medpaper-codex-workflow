"""Version-bound, read-only independent reviews followed by scoped executor fixes.

These records are workflow evidence, not a security sandbox or proof of independence.
The executor must actually delegate; a main-agent self-review is never independent.
"""
from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from pathlib import Path

from .packagefreeze import _safe_path, _sha256, expected_files
from .readiness import atomic_json, reader_inputs

SCOPES = {
    "manuscript": ("M", "manuscript_human_reviewed", "NO_FURTHER_REVIEW", "07_manuscript/human_review.md"),
    "package": ("P", "submission_package_user_confirmed", "OK", "08_submission/package_human_review.md"),
}
ROOT = ".wf/review_rounds"


def now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="microseconds")


def inputs(project: Path, scope: str) -> dict[str, str]:
    if scope == "manuscript":
        files = reader_inputs(project)
        for rel in ("06_refs/refs.bib", "06_refs/refs.ris"):
            if (project / rel).is_file():
                files[rel] = _sha256(project / rel)
        return dict(sorted(files.items()))
    if scope != "package":
        raise ValueError("unknown review scope")
    files = expected_files(project)
    for rel in ("08_submission/bundle/manifest.json", "08_submission/target_journal.json",
                "08_submission/guidelines_extract.md", "08_submission/submission_requirements.json",
                "08_submission/docx_style.json"):
        path = project / rel
        if path.is_file():
            files[rel] = path
    return {rel: _sha256(path) for rel, path in sorted(files.items())}


def latest(project: Path, scope: str) -> dict | None:
    path = project / ROOT / f"{scope}.json"
    return json.loads(path.read_text(encoding="utf-8")) if path.is_file() else None


def save(project: Path, record: dict) -> None:
    atomic_json(project / ROOT / f"{record['id']}.json", record)
    atomic_json(project / ROOT / f"{record['scope']}.json", record)


def changed(before: dict, after: dict) -> list[str]:
    return sorted(rel for rel in set(before) | set(after) if before.get(rel) != after.get(rel))


def start(project: Path, scope: str, *, delta: bool = False) -> dict:
    previous = latest(project, scope)
    if previous and (not previous.get("closed_at") or not previous.get("user_decision")):
        raise ValueError("finish this review round and obtain its user answer before starting another")
    current = inputs(project, scope)
    selection = dict(current)
    differences = changed(previous.get("after", {}), current) if previous else list(current)
    mode = "full"
    if delta and differences and previous and previous.get("report") and (previous.get("user_decision") or {}).get("value") != "abandoned":
        # Deletions are retained explicitly. Unknown dependencies widen to a full review.
        from .dependencies import closure
        try:
            affected = set(closure(project, differences)) | set(differences)
            selection = {rel: digest for rel, digest in current.items() if rel in affected}
            mode = "delta"
        except (OSError, ValueError, TypeError, KeyError):
            mode = "full (dependency uncertainty)"
    number = int(previous["id"][1:]) + 1 if previous else 1
    record = {"schema_version": 1, "id": f"{SCOPES[scope][0]}{number:03d}", "scope": scope,
              "started_at": now(), "mode": mode, "before": current,
              "reviewed_artifacts": selection, "changed_since_previous": differences,
              "deleted_since_previous": [rel for rel in differences if rel not in current],
              "previous_round": previous["id"] if previous else None,
              "report": None, "closed_at": None, "after": {}, "changes": [], "kept": [],
              "user_decision": None}
    save(project, record)
    # Old approval cannot silently release a newly opened round, even with identical files.
    from .state import State
    if (project / ".wf/state.json").is_file():
        state = State(project).load()
        state.data.get("decisions", {}).pop(SCOPES[scope][1], None)
        state.save()
    return record


def restart(project: Path, scope: str, why: str) -> dict:
    record = latest(project, scope)
    if not record or len(why.strip()) < 8:
        raise ValueError("restart requires an existing round and a specific recovery reason")
    record["closed_at"] = record.get("closed_at") or now()
    record["user_decision"] = {"value": "abandoned", "words": why, "at": now()}
    record["recovery_reason"] = why
    save(project, record)
    return start(project, scope)


def attach(project: Path, scope: str, report: Path, *, revalidation: bool = False) -> dict:
    record = latest(project, scope)
    if not record or record.get("closed_at") or (record.get("report") and not revalidation):
        raise ValueError("an open round without an attached report is required")
    current = inputs(project, scope)
    if not revalidation and current != record["before"]:
        raise ValueError("inputs changed while the independent reviewer was reading; restore or restart explicitly")
    if not revalidation:
        for copy in record.get("reader_copies", {}).values():
            if _sha256(_safe_path(project, copy["path"])[1]) != copy["sha256"]:
                raise ValueError("reviewer text copy changed during reading; regenerate and restart this review")
    text = report.read_text(encoding="utf-8")
    metadata = [json.loads(block) for block in re.findall(r"```json\s*\n(.*?)\n```", text, re.S)]
    evidence = [item for item in metadata if isinstance(item, dict) and "reviewed_artifacts" in item]
    selection = current if revalidation else record["reviewed_artifacts"]
    if len(evidence) != 1 or evidence[0]["reviewed_artifacts"] != selection:
        raise ValueError("review report must bind the exact selected input hashes")
    headings = ("Reader comprehension", "Medical presentation") if scope == "manuscript" else (
        "Reader comprehension", "Editorial screening and low-level errors", "Guideline compliance")
    for heading in headings:
        match = re.search(r"(?mi)^#{1,3}\s+" + re.escape(heading) + r"\s*$\n(.*?)(?=^#{1,3}\s|\Z)", text, re.S | re.M)
        if not match or len(match.group(1).strip()) < 80:
            raise ValueError(f"review needs substantive {heading} findings")
    history = record.setdefault("reports", [])
    relative = f"{ROOT}/{record['id']}.review{len(history) + 1}.md"
    destination = project / relative
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(text, encoding="utf-8")
    record["report"] = {"path": relative, "sha256": _sha256(destination), "attached_at": now()}
    history.append(record["report"])
    if revalidation:
        record["reviewed_after"] = current
        record["reviewed_artifacts"] = selection
    save(project, record)
    return record


def close(project: Path, scope: str, changes: list, kept: list, unresolved: int = 0) -> dict:
    record = latest(project, scope)
    if not record or record.get("closed_at") or not record.get("report"):
        raise ValueError("attach the actual independent report before closing an open round")
    report = record["report"]
    if _sha256(project / report["path"]) != report["sha256"]:
        raise ValueError("independent report drifted")
    state_path = project / ".wf/state.json"
    if state_path.is_file() and json.loads(state_path.read_text(encoding="utf-8")).get("active_revision_round"):
        raise ValueError("complete the scoped source revision and its real checks first")
    if unresolved < 0:
        raise ValueError("unresolved defect count cannot be negative")
    for entries in (changes, kept):
        if any(len(entry) != 2 or not all(str(value).strip() for value in entry) for entry in entries):
            raise ValueError("every change or unadopted finding needs a location/ID and concrete reason")
    current = inputs(project, scope)
    missing = set(changed(record["before"], current)) - {entry[0] for entry in changes}
    if missing:
        raise ValueError("record each changed/removed review artifact: " + ", ".join(sorted(missing)))
    record.update(after=current, closed_at=now(), changes=changes, kept=kept, unresolved=unresolved)
    save(project, record)
    summary(project, record)
    return record


def decide(project: Path, scope: str, answer: str, words: str) -> dict:
    record = latest(project, scope)
    if not record or not record.get("closed_at") or record.get("user_decision"):
        raise ValueError("a closed round awaiting the actual user answer is required")
    if answer not in {"another_round", "no_further_review"} or len(words.strip()) < 2:
        raise ValueError("preserve the explicit user answer; generic continuation is not approval")
    if answer == "no_further_review" and words.strip().strip("。.!！").casefold() in {
            "继续", "好的", "好", "ok", "okay", "thanks", "谢谢", "continue", "looks fine"}:
        raise ValueError("generic continuation/acknowledgment is not explicit no-further-review approval")
    if inputs(project, scope) != record["after"]:
        raise ValueError("review files/context changed after closure; do not reuse this approval")
    if answer == "no_further_review" and record.get("unresolved"):
        raise ValueError("unresolved objective defects/user-owned facts prevent approval; preferences may be kept with reasons")
    record["user_decision"] = {"value": answer, "words": words, "at": now()}
    save(project, record)
    summary(project, record)
    if answer == "no_further_review":
        from .state import State
        state = State(project).load()
        state.record_decision(SCOPES[scope][1], SCOPES[scope][2], f"{record['id']}: {words}")
    return record


def summary(project: Path, record: dict) -> None:
    scope = record["scope"]
    headings = ("Materials presented", "Independent verdict", "User-requested revisions", "Revalidation", "Approval to proceed") if scope == "manuscript" else (
        "Package presented", "User modifications", "Revalidation", "User confirmation", "Frozen package")
    lines = [f"# {scope.capitalize()} review {record['id']}", ""]
    content = ["\n".join(f"- `{rel}`" for rel in record["after"]),
               f"Independent report: `{record['report']['path']}`; mode: {record['mode']}",
               "\n".join(f"- Changed {rel}: {why}" for rel, why in record["changes"]) or "No source correction required.",
               "\n".join(f"- Not adopted {key}: {why}" for key, why in record["kept"]) or "No unadopted findings.",
               json.dumps(record.get("user_decision"), ensure_ascii=False)]
    if scope == "package":
        content = [content[0], content[2] + "\n" + content[3],
                   f"Independent report: `{record['report']['path']}`; unresolved defects: {record.get('unresolved', 0)}",
                   content[4], "Freeze follows explicit approval; verify package_review_freeze.json before final audit."]
    # This is the concise current handoff. Complete prior rounds remain in .wf, not chat.
    for heading, body in zip(headings, content):
        lines += [f"## {heading}", "", body, ""]
    (project / SCOPES[scope][3]).write_text("\n".join(lines), encoding="utf-8")


def approved(project: Path, scope: str) -> dict:
    record = latest(project, scope)
    if not record or not record.get("closed_at") or not record.get("report") or record.get("unresolved"):
        raise ValueError("no complete independent review round")
    if (record.get("user_decision") or {}).get("value") != "no_further_review":
        raise ValueError("the user has not explicitly ended this review round")
    if inputs(project, scope) != record["after"]:
        raise ValueError("approved review inputs/context drifted")
    if json.loads((project / ROOT / f"{record['id']}.json").read_text(encoding="utf-8")) != record:
        raise ValueError("review index and sealed round disagree")
    for report in record.get("reports", [record["report"]]):
        if _sha256(_safe_path(project, report["path"])[1]) != report["sha256"]:
            raise ValueError("review report/history drifted")
    return record


def reader_copies(project: Path, files: dict[str, str], folder: str) -> dict:
    """Full local extracts, not substitutes for inspecting actual figures and page images."""
    from .packagecontent import visible_text
    result = {}
    for rel, digest in files.items():
        path = _safe_path(project, rel)[1]
        if _sha256(path) != digest:
            raise ValueError(f"source drifted before extraction: {rel}")
        suffix = path.suffix.lower()
        if suffix == ".docx":
            content = visible_text(path)
            import zipfile
            from xml.etree import ElementTree as ET
            with zipfile.ZipFile(path) as archive:
                for name in sorted(archive.namelist()):
                    if re.fullmatch(r"word/(header\d+|footer\d+|footnotes|endnotes)\.xml", name):
                        root = ET.fromstring(archive.read(name))
                        paragraphs = ["".join(node.text or "" for node in paragraph.iter(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t")) for paragraph in root.iter(
                            "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p")]
                        content += "\nPart: " + name + "\n" + "\n".join(paragraphs)
        elif suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                content = "\n\n".join(f"Sheet: {sheet.title}\n" + "\n".join(
                    "\t".join("" if value is None else str(value) for value in row)
                    for row in sheet.iter_rows(values_only=True)) for sheet in workbook)
            finally:
                workbook.close()
        else:
            continue
        if _sha256(path) != digest:
            raise ValueError(f"source changed during extraction: {rel}")
        import hashlib
        name = hashlib.sha256(rel.encode()).hexdigest()[:16] + ".txt"
        relative = f"{folder}/{name}"
        destination = _safe_path(project, relative)[1]
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_text(f"Source: {rel}\nSHA-256: {digest}\n\n{content}\n", encoding="utf-8")
        result[rel] = {"path": relative, "source_sha256": digest, "sha256": _sha256(destination)}
    return result
