"""Reader-facing diagnostics and explicit medical meaning checks, never a word ban.

No source rewriting, medical inference from keywords, or claim to inspect unprovided papers.
Scanned text is local; a missing text layer still requires actual visual/semantic review.
"""
import hashlib
import json
from pathlib import Path
import re
from .packagefreeze import _safe_path, _sha256
from .readiness import load

RULES = {
    "assistant_voice": (r"\b(?:as an AI(?: language)? model|as requested by (?:the |our )?user|I (?:can|will) (?:help|generate|provide)|here (?:is|are) (?:the |your )?(?:requested |revised )?(?:manuscript|analysis|table|figure)|let(?:'s| us) (?:delve|explore))\b", "Remove executor-to-user conversation from the medical article; retain scientific content and uncertainty, not assistant narration."),
    "workflow_metadata": (r"\b(?:source_sha256|render_sha256|rebuild_files|validation_stages|citekeys_resolve|schema_version|engine_pin|medical_style_version)\b|\bS(?:0[1-9]|1[0-9]|2[0-5])_(?:assemble|package|polish|artifacts|human_review|analysis)\b", "Internal engine fields and stage receipts do not explain clinical methods/results. Keep their provenance backstage; review any genuine scientific software terminology in context."),
    "local_artifact_path": (r"(?:\b[A-Za-z]:[\\/](?:workspace|Users)[\\/]|\b(?:project[\\/])?(?:03_analysis[\\/]results|07_manuscript|08_submission[\\/](?:bundle|integration))[\\/])", "A local workflow path is not a clinical explanation or a stable reader-accessible data source. Replace only the presentation, preserving the actual method and authorized availability statement."),
    "draft_placeholder": (r"\b(?:TODO|TBD|INSERT[_ ](?:RESULT|REFERENCE|FIGURE|TABLE)|lorem ipsum)\b", "Resolve accidental drafting placeholders before delivery; a target-required placeholder or a quoted scientific term needs its actual contextual reason."),
    "internal_plan_history": (r"\b(?:unregistered protocol|archived.{0,15}protocol|protocol.{0,25}(?:archived|not registered|unregistered)|departures from (?:the |our )?protocol|protocol deviations? (?:are|were) (?:reported|listed)|initial (?:analysis )?plan (?:changed|was revised))\b", "Ordinary retrospective mining does not need internal plan/registration history in the paper. Keep the backstage log; preserve actual exploratory/post-hoc status and any genuinely required protocol reporting."),
    "revision_history": (r"\b(?:after internal review|added after review|review round\s+[MP]\d+|after rerunning the pipeline)\b", "Keep orchestration/revision history backstage, but preserve actual post-hoc status and scientific timing."),
    "runtime_count": (r"\b[\d,]+\s+(?:[A-Za-z-]+\s+){0,3}lines?\s+read\b", "Reading-log count: establish clinical counting unit; never relabel file lines as patients/reports."),
    "implementation_detail": (r"\b(?:Unicode normalization|concatenated legacy records|first\s+\d+\s+alphanumeric characters|vocabulary frozen)\b", "Assess scientific necessity in main Methods; move implementation detail, not validity-critical definitions, to supplement when appropriate."),
    "processing_label": (r"\b(?:collapse step|strict(?:\s+[A-Za-z-]+){0,4}\s+contrasts?|publications collapsed|near-duplicates and clusters removed)\b", "Replace/define internal processing labels in clinical terms without changing the procedure or source-specific model."),
    "opaque_endpoint": (r"\bComposite\s+[A-Z]\b", "Use the clinical endpoint name or decode components at the display; a letter is not a clinical definition."),
    "opaque_procedure": (r"\b(?:same-transmitter screen|first-author keys|first receipt|original rule|second rule|record review)\b", "Explain the actual clinical purpose, source unit and population at point of use; a coined processing label is not an explanation. Update the shared clinical vocabulary when appropriate."),
    "ambiguous_antecedent": (r"\bpost\s+hoc\s*:\s*same\b", "Describe the actual post-hoc population/comparison; preserve its timing, not an unexplained 'same'."),
    "software_state": (r"\b(?:NaN|Undefined|Not fitted)\b", "Check missing/not applicable/not fitted/not estimable/unbounded separately; never automatically replace all states with a dash."),
}
FIELDS = {
    "table": ("title_and_labels", "denominators_and_summaries", "footnotes_and_model_states"),
    "figure": ("title_and_labels", "statistic_and_encodings", "population_and_denominators"),
}


def enabled(project, layer="science", final=False):
    from .state import State
    state=State(project).load()
    owner="S17_assemble" if layer=="science" else "S23_package" if final else "S22_polish"
    return not state.pin_excludes(owner,"1.15.0")


def visible(path):
    suffix=path.suffix.lower()
    if suffix in {".md",".txt",".csv"}: return path.read_text(encoding="utf-8")
    if suffix==".docx":
        from .packagecontent import visible_text
        return visible_text(path)
    if suffix==".xlsx":
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        try:
            return "\n".join(f"{sheet.title}!{index}: "+" | ".join(str(c.value) for c in row if c.value is not None)
                             for sheet in book for index,row in enumerate(sheet.iter_rows(),1) if any(c.value is not None for c in row))
        finally: book.close()
    if suffix==".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix==".svg":
        import xml.etree.ElementTree as ET
        root=ET.fromstring(path.read_bytes())
        return "\n".join("".join(node.itertext()) for node in root.iter() if node.tag.endswith("}text") or node.tag=="text")
    return ""  # no OCR inference or substituted code labels for raster figures


def _reader_lines(text):
    """Skip bibliographic titles, retaining locators and later legends/supplements."""
    in_references=False
    for line,value in enumerate(text.splitlines(),1):
        if re.fullmatch(r"\s*(?:#{1,6}\s*)?(?:References|Bibliography)\s*[:.]?\s*",value,re.I):
            in_references=True; continue
        if in_references:
            following=(re.match(r"\s*#{1,6}\s+",value) or
                       re.fullmatch(r"\s*(?:Figure legends?|Table captions?|Tables?|Figures?|Supplementary (?:Methods|Materials?|Information|Tables?|Figures?|Figure legends?)|Declarations(?: and Statements)?|Statements(?: and Declarations)?)\s*[:.]?\s*",value,re.I) or
                       re.match(r"\s*(?:Figure|Table)\s+S?\d+[.:]\s+",value,re.I))
            if not following: continue
            in_references=False
        yield line,value


def _summary_surfaces(text,rel):
    # A count note in Table 2 must not be judged against a median in Table 1 or prose.
    suffix=Path(rel).suffix.lower()
    if suffix==".xlsx":
        sheets={}
        for line in text.splitlines():
            match=re.match(r"^(.*?)!\d+: (.*)$",line)
            if match: sheets.setdefault(match[1],[]).append(match[2])
        return [(f"worksheet {name}","\n".join(lines)) for name,lines in sheets.items()]
    if suffix==".csv": return [("table summary footnote",text)]
    return []  # No reliable table boundary in flattened DOCX/PDF/manuscript text.


def diagnose(text, rel):
    findings=[]
    for line,value in _reader_lines(text):
        for kind,(pattern,message) in RULES.items():
            match=re.search(pattern,value,re.I)
            if match:
                # Hash the complete context, not the display-truncated excerpt. A change
                # late in a long paragraph must not inherit an obsolete retained reason.
                context=value.strip()
                identity=hashlib.sha256(f"{rel}\0{kind}\0{context}".encode()).hexdigest()[:16]
                start=max(0,match.start()-60); end=min(len(value),start+240)
                excerpt=("…" if start else "")+value[start:end].strip()+("…" if end<len(value) else "")
                findings.append({"id":identity,"path":rel,"locator":f"extracted line {line}","kind":kind,"excerpt":excerpt,"message":message})
    for locator,surface in _summary_surfaces(text,rel):
        if re.search(r"median|IQR|interquartile",surface,re.I) and re.search(r"values are (?:numbers|counts) of",surface,re.I) and not re.search(r"unless otherwise|except",surface,re.I):
            identity=hashlib.sha256(f"{rel}\0summary_footnote\0{locator}\0{surface}".encode()).hexdigest()[:16]
            findings.append({"id":identity,"path":rel,"locator":locator,"kind":"summary_footnote","excerpt":"mixed statistical summaries","message":"Footnote must describe the actual row summaries and available-case denominators, not claim every row is a count."})
    # One retained decision covers repeated identical risk in the same surface; locations
    # remain available without asking for duplicate justifications for the same wording.
    unique={}
    for finding in findings:
        if finding["id"] in unique:
            unique[finding["id"]].setdefault("also_at",[]).append(finding["locator"])
        else: unique[finding["id"]]=finding
    return list(unique.values())


def scan(project,layer="science",final=False):
    from .benchmark import display_source
    plan=load(project,"01_protocol/artifact_plan.json" if layer=="science" else "08_submission/integration/artifact_plan.json")
    folder="07_manuscript" if layer=="science" else "08_submission/integration"
    files=set() if final else {f"{folder}/full_manuscript.md"}
    for name in ("methods.md","results.md","discussion.md","supplementary_methods.md"):
        if not final and (project/folder/name).is_file(): files.add(f"{folder}/{name}")
    for rel in (("05_figures/legends.md","04_tables/table_captions.md") if layer=="science" else
                (f"{folder}/figure_legends.md",f"{folder}/table_captions.md")):
        if not final and (project/rel).is_file(): files.add(rel)
    for group in ("main_figures","supp_figures","main_tables","supp_tables"):
        for entry in plan.get(group,[]): files.add(display_source(project,entry,layer,final))
    if final:
        for row in load(project,"08_submission/bundle/manifest.json")["items"]:
            if row.get("upload",True) and row.get("role") in {"manuscript","tables","figure_legends","supplementary"}: files.add(row["file"])
    issues=[]
    for rel in sorted(files):
        _,file=_safe_path(project,rel)
        if file.is_file(): issues.extend(diagnose(visible(file),rel))
    return issues


def layout_inputs(project, report):
    files={}
    for row in report.get("document_layout",[]):
        rel,file=_safe_path(project,row.get("render",""))
        if not rel.startswith("08_submission/evidence/") or file.suffix.lower() not in {".png",".pdf"} or not file.is_file():
            raise ValueError("Final layout observation must reference an actual existing Word-rendered PNG/PDF under evidence/")
        if file.suffix.lower()==".png":
            from PIL import Image
            with Image.open(file) as image: image.verify()
        else:
            from pypdf import PdfReader
            if not PdfReader(file).pages: raise ValueError("Layout PDF has no actual pages")
        files[rel]=_sha256(file)
    return files


def verify(project, report, layer="science", final=False):
    if not enabled(project,layer,final): return []
    from .benchmark import substantive
    problems=[]
    if report.get("medical_style_version")!=1: problems.append("Complete current medical presentation checks using reference/medical-reader-presentation.md; do not auto-upgrade old evidence")
    plan=load(project,"01_protocol/artifact_plan.json" if layer=="science" else "08_submission/integration/artifact_plan.json")
    observed_checks={}
    for group in ("main_figures","supp_figures","main_tables","supp_tables"):
        kind="table" if "tables" in group else "figure"
        for entry in plan.get(group,[]):
            checks=report.get("display_findings",{}).get(entry["id"],{}).get("medical_checks",{})
            for field in FIELDS[kind]:
                if not substantive(checks.get(field)): problems.append(f"{entry['id']}: describe actual {field}, not generic PASS")
            observations=tuple(re.sub(r"\s+"," ",str(checks.get(field,"")).strip()).casefold() for field in FIELDS[kind])
            if all(observations):
                if len(set(observations))==1:
                    problems.append(f"{entry['id']}: distinct medical questions need actual observations, not one copied sentence for every check")
                key=(kind,observations)
                if key in observed_checks:
                    problems.append(f"{entry['id']} and {observed_checks[key]}: all medical observations are copied; inspect each display's actual labels, measures and denominator")
                observed_checks[key]=entry['id']
    for issue in scan(project,layer,final):
        decision=report.get("style_resolutions",{}).get(issue["id"],{})
        if decision.get("action")!="retained" or not substantive(decision.get("reason")) or not str(decision.get("evidence_location","")).strip():
            problems.append(f"{issue['path']} {issue['locator']}: {issue['kind']}; fix at source or justify its actual medical use against this issue ID")
    if final:
        documents={row["file"] for row in load(project,"08_submission/bundle/manifest.json")["items"]
                   if row.get("upload",True) and row.get("role") in {"manuscript","tables","figure_legends","supplementary"}
                   and Path(row["file"]).suffix.lower()==".docx"}
        rows=report.get("document_layout",[])
        if len(rows)!=len({r.get("upload") for r in rows}) or {r.get("upload") for r in rows}!=documents:
            problems.append("Inspect each actual figure/table/manuscript/supplement DOCX layout; no assumed coverage of absent supplements")
        renders=layout_inputs(project,report)
        for row in rows:
            _,source=_safe_path(project,row.get("upload",""))
            if (not source.is_file() or row.get("source_sha256")!=_sha256(source)
                    or row.get("render_sha256")!=renders.get(row.get("render"))
                    or not str(row.get("locator","")).strip() or not substantive(row.get("observation"))):
                problems.append("Layout reading must bind current upload/render bytes and concrete page/figure/table observations")
    return problems
