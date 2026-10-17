"""Reader-facing prose, actual numbering and caption consistency."""
import re
from . import Ctx, Result, check


@check("final_copy_verified")
def final_copy_verified(ctx: Ctx) -> Result:
    if (ctx.state.decision("submission_package_independent_audit") or {}).get("value") != "PASS_AUTHOR_FINAL_COPY":
        return Result(True, "final_copy_verified", "normal independent final-audit path; no author-copy handoff requested")
    from ..finalcopy import verified
    problems = verified(ctx.project)
    return Result(not problems, "final_copy_verified", "; ".join(problems) or "author-copy source/inventory/context bound; not a new independent reading")


@check("prose_not_lists")
def prose_not_lists(ctx: Ctx) -> Result:
    rel = ctx.spec.get("path", "07_manuscript/full_manuscript.md")
    if not ctx.p(rel).is_file(): return Result(False, "prose_not_lists", f"missing {rel}")
    text = re.sub(r"```.*?```", "", ctx.read(rel), flags=re.S)
    sections = re.split(r"(?m)^#\s+", text)
    if "full_manuscript" not in rel: sections = [text]
    problems = []
    for section in sections:
        heading = section.splitlines()[0].strip().lstrip("# ").casefold() if section.strip() else ""
        if "full_manuscript" in rel and heading not in {"introduction", "methods", "results", "discussion"}: continue
        items = re.findall(r"(?m)^\s{0,3}(?:[-*+]\s+|\d+[.)]\s+)\S", section)
        exception = ctx.spec.get("allow_lists_reason")
        if len(items) >= 2 and not exception:
            problems.append(f"{heading or rel}: body is list-led ({len(items)} items); write medical narrative or justify a comparator/journal-supported exception")
    return Result(not problems, "prose_not_lists", "; ".join(problems) or "medical body prose is continuous; display footnotes/declarations excluded")


@check("journal_profile_valid")
def journal_profile_valid(ctx: Ctx) -> Result:
    from ..journalspec import load, aliases, number_text, spelling
    try:
        target = load(ctx.project); aliases(target); number_text("n = 10000", target.get("number_format", {}))
        spelling("centre", target.get("spelling", ""))
        keywords = target.get("keywords", {})
        if not isinstance(keywords, dict): raise ValueError("keywords must be structured")
        if any(type(keywords[key]) is not int or keywords[key] < 1 for key in ("min", "max") if keywords.get(key) is not None): raise ValueError("keyword bounds must be positive integers or null")
        if keywords.get("min") and keywords.get("max") and keywords["min"] > keywords["max"]: raise ValueError("keyword min exceeds max")
        if keywords.get("separator", ", ") not in {",", ", ", ";", "; "}: raise ValueError("invalid keyword separator")
        if "blinded_main_manuscript" in target and type(target["blinded_main_manuscript"]) is not bool: raise ValueError("blinding must be boolean")
        return Result(True, "journal_profile_valid", "typed journal presentation fields validated before expensive builds")
    except (OSError, ValueError, TypeError) as exc: return Result(False, "journal_profile_valid", str(exc))


@check("display_numbering_consistent")
def display_numbering_consistent(ctx: Ctx) -> Result:
    plan = ctx.read_json("01_protocol/artifact_plan.json")
    text = "\n".join(ctx.read(rel) for rel in ctx.spec.get("paths", ["07_manuscript/full_manuscript.md"]) if ctx.p(rel).is_file())
    text = re.split(r"(?mi)^#\s+(?:References|Figure legends)\s*$", text)[0]
    pattern = re.compile(r"\b(Figure|Fig\.?|Table)\s+(S?\d+)([A-Z])?\b", re.I)
    calls = list(pattern.finditer(text)); problems = []
    for kind, supplementary in (("figure", False), ("figure", True), ("table", False), ("table", True)):
        first = []
        for match in calls:
            if match[1].casefold().startswith(kind[:3]) and match[2].upper().startswith("S") == supplementary:
                number = int(match[2].lstrip("Ss"))
                if number not in first: first.append(number)
        if first != sorted(first): problems.append(f"{'supplementary ' if supplementary else ''}{kind}s not in first-citation order: {first}")
    for group in ("main_tables", "supp_tables"):
        for entry in plan.get(group, []):
            path = ctx.p(entry["file"])
            if not path.is_file(): continue
            from openpyxl import load_workbook
            workbook = load_workbook(path, read_only=True, data_only=False)
            try:
                titles = [str(sheet.cell(1, 1).value or "").strip() for sheet in workbook]
            finally: workbook.close()
            caption_path = ctx.p("04_tables/table_captions.md")
            if caption_path.is_file():
                captions = caption_path.read_text(encoding="utf-8")
                candidates = [line.strip().lstrip("# ") for line in captions.splitlines() if re.match(r"^#*\s*" + re.escape(entry["id"]) + r"[.:\s]", line)]
                matching = [title for title in titles if re.match(re.escape(entry["id"]) + r"[.:\s]", title)]
                if candidates and (len(matching) != 1 or not any(value == matching[0] for value in candidates)):
                    problems.append(f"{entry['id']}: caption differs from the actual workbook title")
    for group in ("main_figures", "supp_figures"):
        for entry in plan.get(group, []):
            figure_calls = [match for match in calls if match[1].casefold().startswith("fig") and f"Figure {match[2].upper()}" == entry["id"]]
            panels = entry.get("panels", [])
            if isinstance(panels, list):
                expected = {str(panel.get("label", "") if isinstance(panel, dict) else panel).strip("() ").upper() for panel in panels}
                expected = {label for label in expected if re.fullmatch("[A-Z]", label)}
                observed = {match[3].upper() for match in figure_calls if match[3]}
                if observed and expected and not any(not match[3] for match in figure_calls) and not expected <= observed:
                    problems.append(f"{entry['id']}: missing panel callouts {', '.join(sorted(expected-observed))}")
    return Result(not problems, "display_numbering_consistent", "; ".join(problems) or "first-citation numbering, declared panels and actual table titles agree")
