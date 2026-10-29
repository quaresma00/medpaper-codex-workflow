"""Synthetic interface/negative controls only; never review the user's absent manuscript."""
import json
from pathlib import Path
import tempfile
import unittest
from wfcore import medicalstyle as m
from wfcore.readiness import atomic_json
from wfcore.packagefreeze import _sha256
from wfcore.state import State

class MedicalPresentationTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix="medpaper-medical-style-"); self.addCleanup(self.temp.cleanup)
        self.project=Path(self.temp.name); self.state=State(self.project).create("medpaper","1.15.0","S17_assemble")
        self.plan={"main_tables":[],"main_figures":[],"supp_tables":[],"supp_figures":[]}
        self.data("01_protocol/artifact_plan.json",self.plan)
        self.text("07_manuscript/full_manuscript.md","# Synthetic medical manuscript\n\nNo real findings are represented here.")
    def text(self,rel,value):
        path=self.project/rel; path.parent.mkdir(parents=True,exist_ok=True); path.write_text(value,encoding="utf-8"); return path
    def data(self,rel,value): atomic_json(self.project/rel,value)
    def report(self): return {"medical_style_version":1,"display_findings":{}}
    def test_ordinary_medical_statistical_terms_are_not_a_blacklist(self):
        text="Firth penalized logistic regression, exact confidence intervals, common support, deduplication, Missing, Python 3.13 and R 4.5.2. A validated bioinformatics pipeline and prediction algorithm were assessed by external validation. None of the null hypotheses were rejected."
        self.assertEqual(m.diagnose(text,"methods.md"),[])
    def test_executor_voice_and_internal_paths_are_reviewed_without_rewriting(self):
        text="Here is the revised manuscript. The source_sha256 is stored at C:\\workspace\\study\\results.json."
        path=self.text("07_manuscript/full_manuscript.md",text)
        self.assertTrue({"assistant_voice","workflow_metadata","local_artifact_path"}.issubset({row["kind"] for row in m.scan(self.project)}))
        self.assertEqual(path.read_text(encoding="utf-8"),text)
    def test_references_are_not_author_voice_but_later_legends_are_reviewed(self):
        text="# Results\nClinical outcome.\n# References\n1. As an AI language model: evaluation of clinical answers.\n# Figure legends\nFigure 1. Here is the revised figure."
        issues=m.diagnose(text,"manuscript.md")
        self.assertEqual(len(issues),1)
        self.assertEqual(issues[0]["locator"],"extracted line 6")
        plain=text.replace("# ","")
        self.assertEqual(len(m.diagnose(plain,"manuscript.docx")),1)
    def test_diagnostic_identity_binds_untruncated_paragraph(self):
        prefix="Clinical comparison description. "*12
        middle=" Definition depends on the exact endpoint and source population."*8
        before=m.diagnose(prefix+"Composite A."+middle+" Endpoint one.","methods.md")[0]
        after=m.diagnose(prefix+"Composite A."+middle+" Endpoint two.","methods.md")[0]
        self.assertEqual(before["excerpt"],after["excerpt"])
        self.assertIn("Composite A",before["excerpt"])
        self.assertLessEqual(len(before["excerpt"]),242)
        self.assertNotEqual(before["id"],after["id"])
    def test_repeated_same_risk_needs_one_resolution_with_all_locations(self):
        issues=m.diagnose("Composite A and Composite B.\nComposite A and Composite B.","methods.md")
        self.assertEqual(len(issues),1)
        self.assertEqual(issues[0]["also_at"],["extracted line 2"])
    def test_generic_parser_count_is_not_a_patient_count(self):
        issues=m.diagnose("1,234 demographic lines read","FigureX.svg")
        self.assertEqual(issues[0]["kind"],"runtime_count")
        self.assertIn("counting unit",issues[0]["message"])
    def test_implementation_flag_does_not_rewrite_or_prespecify(self):
        text="A data-derived vocabulary frozen after processing used Unicode normalization."
        self.text("07_manuscript/full_manuscript.md",text)
        before=(self.project/"07_manuscript/full_manuscript.md").read_bytes()
        self.assertTrue(m.scan(self.project)); self.assertEqual((self.project/"07_manuscript/full_manuscript.md").read_bytes(),before)
        self.assertNotIn("predefined",m.visible(self.project/"07_manuscript/full_manuscript.md"))
    def test_actual_xlsx_all_sheets_and_mixed_footnotes(self):
        from openpyxl import Workbook
        path=self.project/"04_tables/t.xlsx"; path.parent.mkdir(parents=True)
        book=Workbook(); sheet=book.active; sheet.title="Main"
        sheet.append(["Median age (IQR)","50 (40-60)"]); sheet.append([None,"Collapse step"])
        sheet.append(["Values are numbers of reports (column percentage)."])
        book.create_sheet("Second").append(["Not fitted"]); book.save(path)
        text=m.visible(path); issues=m.diagnose(text,"04_tables/t.xlsx")
        self.assertIn("Second!1",text)
        self.assertTrue({"summary_footnote","processing_label","software_state"}.issubset({r["kind"] for r in issues}))
    def test_mixed_summary_exception_is_not_falsely_rejected(self):
        self.assertFalse(any(r["kind"]=="summary_footnote" for r in m.diagnose("Main!1: Median (IQR).\nMain!2: Values are numbers of reports unless otherwise indicated.","table.xlsx")))
    def test_summary_rule_does_not_mix_tables_or_prose_with_table_notes(self):
        separate="Characteristics!1: Median age (IQR).\nAdverse events!1: Values are numbers of reports."
        self.assertFalse(any(r["kind"]=="summary_footnote" for r in m.diagnose(separate,"tables.xlsx")))
        self.assertFalse(any(r["kind"]=="summary_footnote" for r in m.diagnose("Methods: age is reported as a median.\nTable 2: Values are numbers of reports.","manuscript.docx")))
        combined="Characteristics!1: Median age (IQR).\nCharacteristics!2: Values are numbers of reports."
        self.assertTrue(any(r["kind"]=="summary_footnote" for r in m.diagnose(combined,"tables.xlsx")))
    def test_no_automatic_collapse_of_missing_or_unbounded_states(self):
        text="Not fitted; Undefined; NA; unbounded confidence interval; observed zero."
        self.text("07_manuscript/full_manuscript.md",text); m.scan(self.project)
        self.assertEqual((self.project/"07_manuscript/full_manuscript.md").read_text(),text)
    def test_raster_text_is_not_inferred_from_a_filename(self):
        path=self.text("05_figures/Figure1.png","not actually an image; scan does not pretend to inspect raster contents")
        self.assertEqual(m.visible(path),"")
    def test_svg_reader_reads_visible_text_not_script_variables(self):
        path=self.text("05_figures/a.svg",'<svg xmlns="http://www.w3.org/2000/svg"><text>Clinical outcome</text><metadata>Composite A</metadata></svg>')
        self.assertEqual(m.visible(path),"Clinical outcome")
    def test_present_issue_cannot_be_marked_fixed_or_closed_as_preference(self):
        self.text("07_manuscript/full_manuscript.md","Composite A was reported.")
        issue=m.scan(self.project)[0]; report=self.report()
        report["style_resolutions"]={issue["id"]:{"action":"edited","reason":"Test statement intentionally unchanged for this negative control.","evidence_location":"Synthetic fixture Methods and panel label."}}
        self.assertTrue(m.verify(self.project,report))
        report["style_resolutions"][issue["id"]]["action"]="retained"
        self.assertEqual(m.verify(self.project,report),[])
    def test_changed_issue_invalidates_old_retained_reason(self):
        self.text("07_manuscript/full_manuscript.md","Composite A is a synthetic label.")
        issue=m.scan(self.project)[0]; report=self.report()
        report["style_resolutions"]={issue["id"]:{"action":"retained","reason":"A synthetic label retained solely to exercise the controlled interface.","evidence_location":"The synthetic fixture has no real clinical endpoint."}}
        self.text("07_manuscript/full_manuscript.md","Composite B is now a different synthetic label.")
        self.assertTrue(m.verify(self.project,report))
    def test_unrelated_line_shift_does_not_reopen_identical_risk(self):
        before=m.diagnose("Composite A is a synthetic label.","manuscript.md")[0]
        after=m.diagnose("A new heading.\n\nComposite A is a synthetic label.","manuscript.md")[0]
        self.assertEqual(before["id"],after["id"])
    def test_new_acceptance_requires_actual_medical_checks(self):
        self.plan["main_figures"]=[{"id":"Figure 1","file":"05_figures/a.svg"}]
        self.data("01_protocol/artifact_plan.json",self.plan)
        self.text("05_figures/a.svg",'<svg xmlns="http://www.w3.org/2000/svg"><text>Synthetic outcome</text></svg>')
        self.assertTrue(m.verify(self.project,self.report()))
        report=self.report(); report["display_findings"]={"Figure 1":{"medical_checks":{key:"Synthetic controlled observation specifically about "+key for key in m.FIELDS["figure"]}}}
        self.assertEqual(m.verify(self.project,report),[])
    def test_copying_one_observation_for_every_medical_question_is_not_acceptance(self):
        self.plan["main_figures"]=[{"id":"Figure 1","file":"05_figures/a.svg"}]
        self.data("01_protocol/artifact_plan.json",self.plan)
        self.text("05_figures/a.svg",'<svg xmlns="http://www.w3.org/2000/svg"><text>Clinical comparison</text></svg>')
        report=self.report(); report["display_findings"]={"Figure 1":{"medical_checks":{key:"All medical aspects have been checked and passed." for key in m.FIELDS["figure"]}}}
        self.assertTrue(any("copied" in item for item in m.verify(self.project,report)))
    def test_copying_whole_review_between_displays_is_rejected_but_shared_denominator_allowed(self):
        self.plan["main_figures"]=[{"id":"Figure 1","file":"05_figures/a.svg"},{"id":"Figure 2","file":"05_figures/b.svg"}]
        self.data("01_protocol/artifact_plan.json",self.plan)
        for name in ("a","b"): self.text(f"05_figures/{name}.svg",'<svg xmlns="http://www.w3.org/2000/svg"><text>Clinical comparison</text></svg>')
        checks={key:"Concrete synthetic observation specifically addressing "+key for key in m.FIELDS["figure"]}
        report=self.report(); report["display_findings"]={name:{"medical_checks":dict(checks)} for name in ("Figure 1","Figure 2")}
        self.assertTrue(any("copied" in item for item in m.verify(self.project,report)))
        report["display_findings"]["Figure 2"]["medical_checks"]["title_and_labels"]="The second figure's title identifies the separate validation population."
        self.assertEqual(m.verify(self.project,report),[])
    def test_final_scan_reads_actual_supplementary_docx_not_duplicate_integration_draft(self):
        from docx import Document
        self.data("08_submission/integration/artifact_plan.json",self.plan)
        self.text("08_submission/integration/full_manuscript.md","Here is the revised manuscript.")
        folder=self.project/"08_submission/bundle"; folder.mkdir(parents=True)
        main=Document(); main.add_paragraph("Actual clinical manuscript."); main.save(folder/"manuscript.docx")
        supplement=Document(); supplement.add_paragraph("Results are stored in 03_analysis/results/final.json."); supplement.save(folder/"supplement.docx")
        self.data("08_submission/bundle/manifest.json",{"items":[{"file":"08_submission/bundle/manuscript.docx","role":"manuscript"},{"file":"08_submission/bundle/supplement.docx","role":"supplementary"}]})
        issues=m.scan(self.project,"journal",True)
        self.assertEqual({row["path"] for row in issues},{"08_submission/bundle/supplement.docx"})
    def test_final_scan_reads_actual_pdf_supplement(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
        self.data("08_submission/integration/artifact_plan.json",self.plan)
        writer=PdfWriter(); page=writer.add_blank_page(612,792)
        font=DictionaryObject({NameObject("/Type"):NameObject("/Font"),NameObject("/Subtype"):NameObject("/Type1"),NameObject("/BaseFont"):NameObject("/Helvetica")})
        page[NameObject("/Resources")]=DictionaryObject({NameObject("/Font"):DictionaryObject({NameObject("/F1"):writer._add_object(font)})})
        stream=DecodedStreamObject(); stream.set_data(b"BT /F1 12 Tf 72 720 Td (As an AI language model, I can help.) Tj ET")
        page[NameObject("/Contents")]=writer._add_object(stream)
        path=self.project/"08_submission/bundle/supplement.pdf"; path.parent.mkdir(parents=True); writer.write(path)
        self.data("08_submission/bundle/manifest.json",{"items":[{"file":"08_submission/bundle/supplement.pdf","role":"supplementary"}]})
        self.assertEqual(m.scan(self.project,"journal",True)[0]["kind"],"assistant_voice")
    def test_existing_completed_pinned_stage_does_not_retroactively_reopen(self):
        self.state.data["engine_pin"]="1.14.0"; self.state.data["stages"]["S17_assemble"]={"status":"done"}; self.state.save()
        self.assertEqual(m.verify(self.project,{}),[])
    def test_layout_requires_real_render_and_current_bytes(self):
        from PIL import Image
        from docx import Document
        self.data("08_submission/integration/artifact_plan.json",self.plan)
        self.text("08_submission/integration/full_manuscript.md","Synthetic final source.")
        upload=self.project/"08_submission/bundle/manuscript.docx"; upload.parent.mkdir(parents=True)
        doc=Document(); doc.add_paragraph("Synthetic medical fixture."); doc.save(upload)
        self.data("08_submission/bundle/manifest.json",{"items":[{"file":"08_submission/bundle/manuscript.docx","role":"manuscript"}]})
        render=self.project/"08_submission/evidence/page.png"; render.parent.mkdir(parents=True)
        Image.new("RGB",(40,40),"white").save(render)
        report=self.report(); self.assertTrue(m.verify(self.project,report,"journal",True))
        report["document_layout"]=[{"upload":"08_submission/bundle/manuscript.docx","source_sha256":_sha256(upload),"render":"08_submission/evidence/page.png","render_sha256":_sha256(render),"locator":"Synthetic first-page layout location only.","observation":"Controlled blank-page interface fixture, not a real manuscript visual review."}]
        self.assertEqual(m.verify(self.project,report,"journal",True),[])
        Image.new("RGB",(40,40),"black").save(render)
        self.assertTrue(m.verify(self.project,report,"journal",True))
    def test_fake_png_receipt_is_not_a_render(self):
        self.text("08_submission/evidence/fake.png","Not rendered image bytes")
        with self.assertRaises(OSError): m.layout_inputs(self.project,{"document_layout":[{"render":"08_submission/evidence/fake.png"}]})

if __name__=="__main__": unittest.main(verbosity=2)
