#!/usr/bin/env python3
"""Behavioral regressions for convergence, batching, final-file extraction and releases."""
from __future__ import annotations

import argparse
import contextlib
import io
import json
import os
from pathlib import Path
import stat
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from wfcore import registry
from wfcore.dependencies import closure, guard_build
from wfcore.packagefreeze import build_freeze, write_freeze, verify_freeze, _sha256, sync_evidence
from wfcore.readiness import CONTRACT, FACTS, CONTRACT_KEYS, STORY_KEYS, atomic_json, freeze_writing, verify_writing, verify_displays
from wfcore.state import State
from wfcore.submission import write_portal, verify_portal, docx_sections

ROOT = Path(__file__).resolve().parents[1]


class OptimizationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="medpaper-optimization-")
        self.project = Path(self.tmp.name) / "project"
        self.project.mkdir()
        self.pipe = registry.load()
        self.state = State(self.project).create("medpaper", self.pipe.meta["version"], "S06_protocol_final")

    def tearDown(self):
        for path in self.project.rglob("*"):
            if path.is_file():
                path.chmod(stat.S_IWRITE | stat.S_IREAD)
        self.tmp.cleanup()

    def text(self, rel, text="Fixture evidence."):
        path = self.project / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def data(self, rel, payload):
        atomic_json(self.project / rel, payload)
        return self.project / rel

    def writing_fixture(self):
        self.data(CONTRACT, {key: "Defined from the real protocol and completed analysis." for key in CONTRACT_KEYS})
        self.data(FACTS, {**{key: "Bounded clinical statement supported by the completed analysis." for key in STORY_KEYS},
                          "result_sources": ["03_analysis/results/main.json"]})
        self.data("03_analysis/results/main.json", {"n": 100, "script": "03_analysis/code/main.py"})
        self.text("03_analysis/code/main.py", "result = 100\n")
        self.text("01_protocol/protocol_final.md")
        self.text("01_protocol/protocol_diff.md")
        self.data("02_data/acquisition_manifest.json", {"scope": "all protocol data"})
        self.state.record_decision("analysis_converged", "YES", "All relevant analyses and validation have completed.")
        self.state.record_decision("go_nogo_2", "GO", "The supported clinical finding has a realistic journal fit.")

    def test_writing_contract_binds_scientific_inputs(self):
        self.writing_fixture()
        self.text("03_analysis/code/helpers/model.py", "# diagnostic model helper")
        freeze_writing(self.project)
        self.assertEqual(verify_writing(self.project), [])
        self.text("03_analysis/code/helpers/model.py", "# changed model helper")
        self.assertTrue(verify_writing(self.project))
        freeze_writing(self.project)
        self.data("03_analysis/results/main.json", {"n": 99})
        self.assertTrue(verify_writing(self.project))

    def test_unresolved_story_and_non_go_cannot_freeze(self):
        self.writing_fixture()
        facts = json.loads((self.project / FACTS).read_text())
        facts["cannot_claim"] = "TODO"
        self.data(FACTS, facts)
        with self.assertRaises(ValueError):
            freeze_writing(self.project)
        self.state.record_decision("go_nogo_2", "STOP", "Not scientifically interpretable; do not write a manuscript.")
        with self.assertRaises(ValueError):
            freeze_writing(self.project)

    def test_display_review_detects_changed_preview_and_missing_coverage(self):
        preview = self.text("01_protocol/prototypes/Table1.md", "| Group | Patients |\n| A | 100 |")
        source = self.data("03_analysis/results/main.json", {"n": 100})
        self.data("01_protocol/artifact_plan.json", {"main_tables": [{"id": "Table 1", "source_results": ["03_analysis/results/main.json"]}]})
        row = {"id": "Table 1", "preview": "01_protocol/prototypes/Table1.md",
               "preview_sha256": _sha256(preview), "source_hashes": {"03_analysis/results/main.json": _sha256(source)},
               "question": "Who was included?", "reader_explanation": "The table describes the full eligible cohort.",
               "denominator_units_missingness": "100 patients, with complete group ascertainment.",
               "encoding_overlap": "Mutually exclusive groups; counts are persons.",
               "layout": "Portrait, two columns at readable type size.", "reviewer": "AI reader simulation", "verdict": "PASS"}
        self.data("01_protocol/display_review.json", {"items": [row]})
        self.assertFalse(verify_displays(self.project))
        preview.write_text("Changed meaning", encoding="utf-8")
        self.assertTrue(verify_displays(self.project))
        self.data("01_protocol/display_review.json", {"items": []})
        self.assertTrue(verify_displays(self.project))

    def command(self, *args):
        env = {**os.environ, "MEDPAPER_PROJECT": str(self.project), "MEDPAPER_ROOT": str(ROOT), "PYTHONIOENCODING": "utf-8"}
        return subprocess.run([sys.executable, str(ROOT / "tools/rework.py"), *args,
                               "--project", str(self.project)], env=env, capture_output=True, text=True, encoding="utf-8")

    def batch_fixture(self):
        self.state.data["current"] = "S24_package_human_review"
        self.state.save()
        plan = self.data("temp/feedback.json", {
            "feedback_verbatim": "Please change the corresponding author address on the title page.",
            "interpretation": "Update the supplied administrative fact and only its actual document consumers.",
            "items": [{"kind": "title-page-or-statements", "request": "Correct the corresponding author's postal address.",
                       "affected_sources": ["00_input/author_info.json"],
                       "acceptance_criteria": ["All actual consumers use the supplied correct address."]}]})
        return plan

    def test_open_batch_prevents_build_until_sealed_and_deduplicates(self):
        plan = self.batch_fixture()
        self.assertEqual(self.command("batch", "--plan", str(plan)).returncode, 0)
        self.assertEqual(State(self.project).load().current, "S24_package_human_review")
        with self.assertRaises(ValueError):
            guard_build(self.project)
        self.assertEqual(self.command("batch", "--plan", str(plan)).returncode, 0)
        row = json.loads((self.project / ".wf/revisions/R001.json").read_text())
        self.assertEqual(len(row["items"]), 1)
        self.assertNotEqual(self.command("build", "--scope", "full").returncode, 0)
        self.assertEqual(self.command("seal", "--why", "The user has finished the batch and asked to apply it now.").returncode, 0)
        self.assertEqual(State(self.project).load().current, "S21_authors")
        guard_build(self.project)

    def test_third_full_build_switches_to_targeted_without_closing_round(self):
        plan = self.batch_fixture()
        self.assertEqual(self.command("batch", "--sealed", "--plan", str(plan)).returncode, 0)
        self.assertEqual(self.command("build", "--scope", "full").returncode, 0)
        self.assertEqual(self.command("build", "--scope", "full").returncode, 0)
        self.assertNotEqual(self.command("build", "--scope", "full").returncode, 0)
        self.assertEqual(self.command("build", "--scope", "targeted", "--why", "Only the changed title page needs correction.").returncode, 0)
        self.assertEqual(State(self.project).load().data["active_revision_round"], "R001")

    def test_apply_now_can_seal_previously_collected_identical_feedback(self):
        plan = self.batch_fixture()
        self.assertEqual(self.command("batch", "--plan", str(plan)).returncode, 0)
        self.assertEqual(self.command("batch", "--sealed", "--plan", str(plan)).returncode, 0)
        row = json.loads((self.project / ".wf/revisions/R001.json").read_text())
        self.assertEqual(len(row["items"]), 1)
        self.assertEqual(row["status"], "active")
        self.assertEqual(State(self.project).load().current, "S21_authors")

    def test_dependency_closure_does_not_rebuild_unrelated_science(self):
        self.data("08_submission/bundle/manifest.json", {"items": [
            {"role": "title_page", "file": "08_submission/bundle/title.docx"},
            {"role": "manuscript", "file": "08_submission/bundle/manuscript.docx"}]})
        changed = closure(self.project, ["08_submission/integration/title_page.md"])
        self.assertIn("08_submission/bundle/title.docx", changed)
        self.assertNotIn("08_submission/bundle/manuscript.docx", changed)
        self.assertNotIn("06_refs/library.json", changed)

    def test_changed_data_and_combined_supplement_keep_their_consumers(self):
        self.text("03_analysis/code/main.py", "# executed analysis")
        self.data("03_analysis/results/main.json", {"built_by": "03_analysis/code/main.py", "n": 100})
        self.data("08_submission/bundle/manifest.json", {"items": [
            {"role": "supplementary", "file": "08_submission/bundle/supplement.docx",
             "source_files": ["08_submission/integration/supplementary_methods.md", "04_tables/supplementary/tables.xlsx"]}]})
        for source in ("02_data/raw/cohort.csv", "01_protocol/analysis_contract.json"):
            changed = closure(self.project, [source])
            self.assertIn("03_analysis/results/main.json", changed)
            self.assertIn("07_manuscript/full_manuscript.md", changed)
        changed = closure(self.project, ["04_tables/supplementary/tables.xlsx"])
        self.assertIn("08_submission/bundle/supplement.docx", changed)
        self.assertNotIn("03_analysis/results/main.json", changed)

    def test_default_status_never_executes_gates(self):
        from wfcore import cli
        with patch.object(cli, "_load", return_value=(self.pipe, self.state, self.project)), \
             patch.object(cli.gates, "run_stage", side_effect=AssertionError("status ran gates")), \
             contextlib.redirect_stdout(io.StringIO()) as output:
            self.assertEqual(cli.cmd_status(argparse.Namespace(full=False, brief=False, json=False)), 0)
        self.assertIn("S06_protocol_final", output.getvalue())
        self.assertNotIn("NON-NEGOTIABLE INVARIANTS", output.getvalue())

    def upload_fixture(self):
        self.text("08_submission/bundle/manuscript.docx", "immutable file bytes")
        self.data("08_submission/bundle/manifest.json", {"items": [{"role": "manuscript", "file": "08_submission/bundle/manuscript.docx"}]})
        self.text("08_submission/guidelines_extract.md", "Official rules for this article type.")
        self.data("08_submission/target_journal.json", {"journal": "Test Journal"})

    def test_evidence_refresh_preserves_upload_freeze(self):
        self.upload_fixture()
        write_freeze(self.project)
        before = build_freeze(self.project)["freeze_id"]
        self.text("08_submission/cache/new.html", "A renewed retrieval receipt.")
        self.text("08_submission/submission_qc.md", "Updated backstage QC note.")
        sync_evidence(self.project)
        self.assertTrue(verify_freeze(self.project)[0])
        self.assertEqual(before, build_freeze(self.project)["freeze_id"])

    def test_release_readonly_and_content_are_verified(self):
        self.upload_fixture()
        freeze = write_freeze(self.project)
        payload = json.loads(freeze.read_text())
        release_file = self.project / payload["release_dir"] / "manuscript.docx"
        self.assertFalse(release_file.stat().st_mode & stat.S_IWUSR)
        release_file.chmod(stat.S_IWRITE | stat.S_IREAD)
        self.assertFalse(verify_freeze(self.project)[0])
        release_file.write_text("tampered", encoding="utf-8")
        release_file.chmod(stat.S_IREAD)
        self.assertFalse(verify_freeze(self.project)[0])
        with self.assertRaises(ValueError):
            write_freeze(self.project)

    def test_release_metadata_cannot_collide_with_an_upload(self):
        self.upload_fixture()
        self.text("08_submission/bundle/upload_manifest.json", "Not release metadata")
        self.data("08_submission/bundle/manifest.json", {"items": [
            {"role": "manuscript", "file": "08_submission/bundle/manuscript.docx"},
            {"role": "supplementary", "file": "08_submission/bundle/upload_manifest.json"}]})
        with self.assertRaises(ValueError):
            write_freeze(self.project)

    def test_upload_change_changes_identity_but_old_release_survives(self):
        self.upload_fixture()
        freeze = write_freeze(self.project)
        old = json.loads(freeze.read_text())
        self.text("08_submission/bundle/manuscript.docx", "user requested changed bytes")
        self.assertFalse(verify_freeze(self.project)[0])
        self.assertNotEqual(old["freeze_id"], build_freeze(self.project)["freeze_id"])
        self.assertEqual((self.project / old["release_dir"] / "manuscript.docx").read_text(), "immutable file bytes")

    def portal_fixture(self):
        from docx import Document
        path = self.project / "08_submission/bundle/manuscript.docx"
        path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()
        for text in ("Clinical cohort findings", "Abstract", "Background", "A clinical question.",
                     "Methods", "We studied adults.", "Results", "Associations were observed.",
                     "Conclusions", "Causation remains uncertain.", "Keywords: cohort studies, prognosis",
                     "Introduction", "Clinical uncertainty remains.", "Methods", "A retrospective cohort.",
                     "Results", "One hundred patients.", "Discussion", "Interpret with caution.", "References", "One reference."):
            doc.add_paragraph(text)
        doc.save(path)
        self.data("08_submission/bundle/manifest.json", {"items": [{"role": "manuscript", "file": "08_submission/bundle/manuscript.docx"}]})
        self.data("08_submission/submission_requirements.json", {"article_type": "Original Article", "required_upload_roles": ["manuscript"],
            "word_count": {"include_sections": ["Introduction", "Methods", "Results", "Discussion"],
                           "limit": 100, "abstract_limit": 100, "source": "Official journal instructions"}})
        self.data("00_input/author_info.json", {key: "User-supplied fact" for key in
            ("authors", "affiliations", "corresponding", "funding", "conflicts", "ethics_approval", "data_availability")})
        return path

    def test_portal_uses_final_docx_and_keeps_structured_abstract_together(self):
        path = self.portal_fixture()
        payload = write_portal(self.project)
        self.assertEqual(payload["word_counts"]["main"], 12)
        self.assertIn("We studied adults.", payload["abstract"])
        self.assertEqual(payload["keywords"], ["cohort studies", "prognosis"])
        self.assertFalse(verify_portal(self.project))
        self.text("07_manuscript/full_manuscript.md", "An unrelated old Markdown draft.")
        self.assertFalse(verify_portal(self.project))
        from docx import Document
        doc = Document(path)
        doc.paragraphs[0].text = "Changed final title"
        doc.save(path)
        self.assertTrue(verify_portal(self.project))

    def test_final_docx_over_limit_is_rejected(self):
        self.portal_fixture()
        requirements = json.loads((self.project / "08_submission/submission_requirements.json").read_text())
        requirements["word_count"]["limit"] = 5
        self.data("08_submission/submission_requirements.json", requirements)
        with self.assertRaises(ValueError):
            write_portal(self.project)

    def test_pmcid_and_nonfirst_author_forgery_are_rejected(self):
        from pubmed import eutils
        from wfcore import refproof
        raw = '<PubmedArticleSet><PubmedArticle><MedlineCitation><PMID>123</PMID><Article><ArticleTitle>Clinical study</ArticleTitle><Abstract><AbstractText>Supported findings.</AbstractText></Abstract><Journal><Title>Clinical Journal</Title><JournalIssue><PubDate><Year>2025</Year></PubDate></JournalIssue></Journal><AuthorList><Author><LastName>van der Meer</LastName><ForeName>Anna</ForeName><Initials>A</Initials></Author><Author><CollectiveName>Study Research Group</CollectiveName></Author></AuthorList></Article></MedlineCitation><PubmedData><ArticleIdList><ArticleId IdType="pmc">PMC123</ArticleId><ArticleId IdType="doi">10.1000/example</ArticleId></ArticleIdList></PubmedData></PubmedArticle></PubmedArticleSet>'
        source = refproof.parse_pubmed_payload(raw)["123"]
        entry = eutils.parse_pubmed_xml(raw)[0]
        self.assertFalse(refproof.compare_entry_to_source(entry, source))
        entry["pmcid"] = "PMC456"
        self.assertIn("PMCID differs from PubMed payload", refproof.compare_entry_to_source(entry, source))
        entry["pmcid"] = "PMC123"
        entry["authors"][1]["last"] = "Invented Group"
        self.assertTrue(any("ordered authors" in p for p in refproof.compare_entry_to_source(entry, source)))

    def test_live_cache_requires_current_identity_hash_and_time(self):
        from datetime import datetime, timedelta, timezone
        from pubmed import eutils
        from wfcore.liverefs import cached_sources, store_sources, REL
        raw = '<PubmedArticleSet><PubmedArticle><MedlineCitation><PMID>123</PMID><Article><ArticleTitle>Fixture study</ArticleTitle><Abstract><AbstractText>Fixture abstract.</AbstractText></Abstract><Journal><Title>Fixture Journal</Title><JournalIssue><PubDate><Year>2025</Year></PubDate></JournalIssue></Journal></Article></MedlineCitation></PubmedArticle></PubmedArticleSet>'
        rel = "06_refs/cache/gate_efetch_pubmed_fixture.xml"
        path = self.text(rel, raw)
        entry = eutils.parse_pubmed_xml(raw)[0]
        entry["cache_file"] = rel
        store_sources(self.project, [entry])
        self.assertIsNotNone(cached_sources(self.project, [entry]))
        entry["pmcid"] = "PMC999"
        self.assertIsNone(cached_sources(self.project, [entry]))
        entry["pmcid"] = ""
        payload = json.loads((self.project / REL).read_text())
        payload["at"] = (datetime.now(timezone.utc) - timedelta(hours=25)).isoformat()
        self.data(REL, payload)
        self.assertIsNone(cached_sources(self.project, [entry]))
        store_sources(self.project, [entry])
        path.write_text(raw + "\n", encoding="utf-8")
        self.assertIsNone(cached_sources(self.project, [entry]))

    def test_s13_and_s25_never_trust_cache_and_fail_closed_offline(self):
        from wfcore.checks import Ctx
        from wfcore.checks.refs import reference_provenance
        self.data("06_refs/library.json", {"entries": [{"pmid": "123", "citekey": "fixture"}]})
        for stage in ("S13_reflib", "S25_submission_audit"):
            ctx = Ctx(self.pipe, self.state, self.project, self.pipe.stage(stage), {"live": True})
            with patch("wfcore.refproof.validate_local_proof", return_value=(True, [], 1)), \
                 patch("wfcore.liverefs.cached_sources", return_value={"123": {}}) as cache, \
                 patch("pubmed.eutils.efetch", side_effect=OSError("offline fixture")) as fetch:
                self.assertFalse(reference_provenance(ctx).ok)
                cache.assert_not_called()
                fetch.assert_called_once_with(["123"], fresh=True, purpose="gate")

    def test_bibtex_preserves_collective_and_particle_names_in_pandoc(self):
        from pubmed.build_library import to_bibtex
        import shutil
        if not shutil.which("pandoc"):
            self.skipTest("Pandoc is optional for this export-reader integration test")
        raw = to_bibtex([{"citekey": "nameFixture", "title": "Clinical study", "year": "2025", "authors": [
            {"last": "van der Meer", "first": "Anna"}, {"last": "Study Research Group", "collective": True}]}])
        parsed = subprocess.run(["pandoc", "-f", "biblatex", "-t", "csljson"], input=raw,
                                capture_output=True, text=True, encoding="utf-8", check=True)
        authors = json.loads(parsed.stdout)[0]["author"]
        self.assertEqual(authors[1]["literal"], "Study Research Group")
        self.assertEqual(authors[0]["family"], "Meer")
        self.assertEqual(authors[0]["non-dropping-particle"], "van der")


if __name__ == "__main__":
    unittest.main(verbosity=2)
