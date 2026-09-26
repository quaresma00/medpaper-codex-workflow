#!/usr/bin/env python3
"""Behavior tests for scoped edits; synthetic fixtures never touch a real study."""
import argparse
import contextlib
import copy
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from wfcore import registry, revision, cli
from wfcore.readiness import atomic_json, reader_inputs, verify_reader_review
from wfcore.state import State
from wfcore.packagecontent import write_baseline, verify_baseline

ROOT = Path(__file__).resolve().parents[1]


class ScopedTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="medpaper-scoped-")
        self.project = Path(self.tmp.name) / "project"
        self.project.mkdir()
        self.pipe = registry.load()
        self.state = State(self.project).create("medpaper", self.pipe.meta["version"], "S19_human_review")
        for stage in self.pipe.stages[:23]:
            self.state.stage_info(stage.id)["status"] = "done"
        self.state.save()

    def tearDown(self):
        self.tmp.cleanup()

    def text(self, rel, content="Synthetic unchanged source"):
        path = self.project / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        return path

    def data(self, rel, value):
        atomic_json(self.project / rel, value)
        return self.project / rel

    def command(self, *args, tool="rework.py"):
        env = {**os.environ, "MEDPAPER_PROJECT": str(self.project), "MEDPAPER_ROOT": str(ROOT), "PYTHONIOENCODING": "utf-8"}
        return subprocess.run([sys.executable, str(ROOT / "tools" / tool), *args], env=env,
                              text=True, encoding="utf-8", capture_output=True)

    def open(self, sources, kind="figure-layout", checkpoint="S19_human_review", extra=None):
        self.state.data["current"] = checkpoint
        self.state.save()
        item = {"kind": kind, "request": "Implement this focused user-requested correction without altering unrelated artifacts.",
                "affected_sources": sources, "acceptance_criteria": ["Only intended outputs change", "Actual checks pass"]}
        plan = self.data("temp/plan.json", {"feedback_verbatim": "Please apply this focused correction while retaining all unaffected work.",
                    "interpretation": "Update the canonical source and only its real consumers.", "items": [item] + (extra or [])})
        result = self.command("batch", "--sealed", "--plan", str(plan))
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.state.load()
        return revision.active(self.project)

    def figures(self):
        rows = []
        for number in (1, 2):
            stem = f"05_figures/out/Figure{number}"
            for suffix in (".py", ".png", ".pdf"):
                self.text(stem + suffix)
            rows.append({"id": f"Figure {number}", "file": stem + ".png", "pdf": stem + ".pdf",
                         "script": stem + ".py", "source_results": ["03_analysis/results/main.json"]})
        self.data("01_protocol/artifact_plan.json", {"main_figures": rows})
        self.data("03_analysis/results/main.json", {"n": 123})
        return rows

    def test_figure_layout_keeps_history_and_protects_other_figure(self):
        self.figures()
        old = copy.deepcopy(self.state.data["stages"])
        row = self.open(["05_figures/out/Figure2.py"])
        self.assertEqual(self.state.current, "S19_human_review")
        self.assertEqual(old, self.state.data["stages"])
        self.assertIn("05_figures/out/Figure2.pdf", row["rebuild_files"])
        self.assertIn("05_figures/qc/Figure2.artist.json", row["rebuild_files"])
        self.assertNotIn("05_figures/out/Figure1.pdf", row["rebuild_files"])
        self.assertNotIn("03_analysis/results/main.json", row["rebuild_files"])
        revision.guard_outputs(self.project, [self.project / "05_figures/out/Figure2.pdf"])
        with self.assertRaises(ValueError):
            revision.guard_outputs(self.project, [self.project / "05_figures/out/Figure1.pdf"])
        self.text("05_figures/out/Figure1.pdf", "Unplanned overwrite")
        self.assertTrue(revision.scope_problems(self.project, row))

    def test_no_force_advance_or_loop_escape(self):
        self.figures()
        self.open(["05_figures/out/Figure2.py"])
        with patch.object(cli, "_load", return_value=(self.pipe, self.state, self.project)), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cli.cmd_advance(argparse.Namespace(force=True, note="Skip")), 2)
            self.assertEqual(cli.cmd_loop(argparse.Namespace(to="S11_figures", why="restart")), 2)
        self.assertEqual(self.state.current, "S19_human_review")

    def test_cosmetic_request_cannot_change_analysis(self):
        self.text("03_analysis/code/model.py")
        plan = self.data("temp/bad.json", {"feedback_verbatim": "Change the model but call this a layout adjustment to save time.",
            "interpretation": "This is an invalid type downgrade fixture.", "items": [{"kind": "figure-layout",
            "affected_sources": ["03_analysis/code/model.py"], "request": "Alter the model as a cosmetic change.",
            "acceptance_criteria": ["This must fail type validation"]}]})
        self.assertNotEqual(self.command("batch", "--sealed", "--plan", str(plan)).returncode, 0)

    def test_late_science_retains_deferred_package_items(self):
        self.text("07_manuscript/discussion.md")
        self.state.record_decision("manuscript_human_reviewed", "NO_FURTHER_REVIEW", "Prior scientific approval")
        row = self.open(["07_manuscript\\discussion.md"], "discussion", "S24_package_human_review", [{
            "kind": "title-page-or-statements", "request": "Correct the author's supplied address on the title page.",
            "affected_sources": ["00_input/author_info.json"], "acceptance_criteria": ["Correct supplied address appears"]}])
        self.assertEqual(self.state.current, "S19_human_review")
        self.assertEqual(len(row["deferred_package_items"]), 1)
        self.assertIsNone(self.state.decision("manuscript_human_reviewed"))
        self.assertTrue(self.state.is_done("S21_authors"))
        self.assertIn("S18_independent_review", row["validation_stages"])
        self.assertNotIn("08_submission/integration/full_manuscript.md", row["rebuild_files"])

    def test_resume_does_not_replay_stage_tail_or_lose_feedback(self):
        self.text("07_manuscript/discussion.md")
        self.text("07_manuscript/full_manuscript.md")
        row = self.open(["07_manuscript/discussion.md"], "discussion", "S24_package_human_review")
        self.text("07_manuscript/discussion.md", "Corrected scientific interpretation")
        self.text("07_manuscript/full_manuscript.md", "Assembled corrected scientific interpretation")
        # Unit-test the transition helper; CLI gating before this helper is tested separately.
        row["status"] = "complete"
        self.data(".wf/revisions/R001.json", row)
        self.state.data.pop("active_revision_round")
        self.state.save()
        round_id = revision.resume_package(self.project, self.pipe, self.state)
        self.assertEqual(round_id, "R002")
        self.assertEqual(self.state.current, "S24_package_human_review")
        self.assertTrue(self.state.is_done("S20_journal"))
        resumed = revision.active(self.project)
        self.assertIn("08_submission/integration/full_manuscript.md", resumed["rebuild_files"])
        self.assertNotIn("07_manuscript/discussion.md", resumed["allowed_files"])

    def test_closed_round_requires_real_checks_not_labels(self):
        source = self.text("08_submission/integration/title_page.md", "# Title page\n\nA correct address.")
        self.open(["08_submission/integration/title_page.md"], "title-page-or-statements", "S24_package_human_review")
        result = self.command("mark", "--item", "R001-01", "--changed-file", "08_submission/integration/title_page.md",
                              "--validated-by", "I claim all passed", "--summary", "Updated the source and reviewed its final administrative wording.")
        self.assertEqual(result.returncode, 0, result.stderr)
        closed = self.command("close", "--summary", "All requested changes were implemented and checked against the acceptance criteria.")
        self.assertNotEqual(closed.returncode, 0)
        checked = self.command("check")
        self.assertNotEqual(checked.returncode, 0)  # A title alone cannot satisfy real journal gates.
        self.assertIn("Scoped checks", checked.stdout)

    def test_scoped_check_actual_validator_and_fingerprint(self):
        self.text("08_submission/integration/title_page.md", "# Title page\n\nA correct address.")
        row = self.open(["08_submission/integration/title_page.md"], "title-page-or-statements", "S24_package_human_review")
        # Narrow unit scope uses the real Markdown validator, not a fabricated PASS result.
        row["checks"] = [{"stage": "S21_authors", "spec": {"check": "md_sections",
                         "path": "08_submission/integration/title_page.md", "headings": ["Title page"]}}]
        path = self.project / ".wf/revisions/R001.json"
        self.data(".wf/revisions/R001.json", row)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(revision.check_round(self.project, self.pipe, self.state, path, row), 0)
        signature = row["validation"]["input_signature"]
        self.assertEqual(signature, revision.input_signature(self.project, row))
        self.text("08_submission/integration/title_page.md", "Missing required heading")
        self.assertNotEqual(signature, revision.input_signature(self.project, row))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(revision.check_round(self.project, self.pipe, self.state, path, row), 2)

    def test_style_only_cannot_recapture_word_text(self):
        from docx import Document
        docx = self.project / "08_submission/bundle/title.docx"
        docx.parent.mkdir(parents=True)
        doc = Document(); doc.add_paragraph("Original title"); doc.save(docx)
        self.text("08_submission/integration/title_page.md", "Original title")
        self.data("08_submission/bundle/manifest.json", {"items": [{"role": "title_page", "file": "08_submission/bundle/title.docx"}]})
        write_baseline(self.project)
        self.open(["08_submission/bundle/title.docx"], "word-format-only", "S24_package_human_review")
        doc.paragraphs[0].text = "Untracked narrative change"; doc.save(docx)
        self.assertFalse(verify_baseline(self.project)[0])
        self.assertFalse(revision.baseline_capture_allowed(self.project))
        self.assertNotEqual(self.command("capture", "--replace", "--project", str(self.project), tool="package_content.py").returncode, 0)

    def test_source_rebuild_receipt_must_match_current_output_and_input(self):
        from docx import Document
        docx = self.project / "08_submission/bundle/title.docx"
        docx.parent.mkdir(parents=True)
        doc = Document(); doc.add_paragraph("Original title"); doc.save(docx)
        source = self.text("08_submission/integration/title_page.md", "Original title")
        self.data("08_submission/bundle/manifest.json", {"items": [{"role": "title_page", "file": "08_submission/bundle/title.docx"}]})
        write_baseline(self.project)
        self.open(["08_submission/integration/title_page.md"], "title-page-or-statements", "S24_package_human_review")
        source.write_text("Revised title", encoding="utf-8")
        doc.paragraphs[0].text = "Revised title"; doc.save(docx)
        self.assertFalse(revision.baseline_capture_allowed(self.project))
        revision.record_build(self.project, docx, [source])
        self.assertTrue(revision.baseline_capture_allowed(self.project))
        source.write_text("Changed again", encoding="utf-8")
        self.assertFalse(revision.baseline_capture_allowed(self.project))

    def test_reader_review_covers_actual_figure_and_detects_drift(self):
        self.figures()
        for rel in ("07_manuscript/full_manuscript.md", "05_figures/legends.md", "04_tables/table_captions.md"):
            self.text(rel)
        metadata = {"reviewed_artifacts": reader_inputs(self.project)}
        self.text("07_manuscript/independent_publishability_review.md", "# Reader comprehension\n\n" +
                  "The cohort and comparator are identifiable, but association does not establish causation. " * 2 +
                  "\n# Medical presentation\n\n" + "Clinical labels and denominators were inspected in the actual rendered displays. " * 2 +
                  "\n```json\n" + json.dumps(metadata) + "\n```\n")
        self.assertFalse(verify_reader_review(self.project))
        self.text("05_figures/out/Figure2.png", "Changed rendered file")
        self.assertTrue(verify_reader_review(self.project))

    def test_missing_independent_review_does_not_get_fabricated(self):
        self.figures()
        self.assertTrue(verify_reader_review(self.project))

    def test_real_validator_can_close_round_and_detect_later_deliverable_drift(self):
        from rework import _cmd_mark, _cmd_close
        from wfcore.checks.revisions import revision_rounds_closed
        from wfcore.checks import Ctx
        rel = "08_submission/integration/title_page.md"
        self.text(rel, "# Title page\n\nCurrent correct address.")
        row = self.open([rel], "title-page-or-statements", "S24_package_human_review")
        # Explicitly narrow unit fixture; production plans retain all owner gates.
        row["checks"] = [{"stage": "S21_authors", "spec": {"check": "md_sections", "path": rel, "headings": ["Title page"]}}]
        self.data(".wf/revisions/R001.json", row)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(revision.check_round(self.project, self.pipe, self.state, self.project / ".wf/revisions/R001.json", row), 0)
            _cmd_mark(argparse.Namespace(item="R001-01", summary="The corrected title source was checked by the actual Markdown validator.",
                                        changed_file=[rel], validated_by=["Actual md_sections check"]), self.project, self.state)
            self.assertEqual(_cmd_close(argparse.Namespace(summary="All fixture changes passed the actual scoped validator and preserved the unaffected files."), self.state), 0)
        ctx = Ctx(self.pipe, self.state, self.project, self.pipe.stage("S24_package_human_review"), {})
        self.assertTrue(revision_rounds_closed(ctx).ok)
        self.text(rel, "Changed after approval")
        self.assertFalse(revision_rounds_closed(ctx).ok)

    def test_rebase_refuses_unmerged_science_and_preserves_journal_edits(self):
        from wfcore import reviewpackage, scientificfreeze, journalworkspace
        for rel in (*reviewpackage.REQUIRED_FILES, "07_manuscript/human_review.md"):
            self.text(rel, "Old scientific content\n")
        self.data("08_submission/target_journal.json", {"journal": "Synthetic Journal", "issn": "0000-0000"})
        reviewpackage.build(self.project, self.pipe.meta["version"])
        scientificfreeze.write(self.project, self.pipe.meta["version"])
        journalworkspace.initialise(self.project)
        self.text("08_submission/integration/full_manuscript.md", "Old scientific content\nJournal-specific retained wording\n")
        row = self.open(["07_manuscript/full_manuscript.md"], "discussion", "S24_package_human_review")
        self.text("07_manuscript/full_manuscript.md", "Corrected scientific content\n")
        row["status"] = "complete"
        self.data(".wf/revisions/R001.json", row)
        self.state.data.pop("active_revision_round")
        self.state.save()
        reviewpackage.build(self.project, self.pipe.meta["version"])
        scientificfreeze.write(self.project, self.pipe.meta["version"])
        revision.resume_package(self.project, self.pipe, self.state)
        with self.assertRaisesRegex(ValueError, "not merged"):
            journalworkspace.rebase(self.project, "Merged all changed scientific passages while retaining journal-specific wording.")
        merged = "Corrected scientific content\nJournal-specific retained wording\n"
        self.text("08_submission/integration/full_manuscript.md", merged)
        journalworkspace.rebase(self.project, "Merged all changed scientific passages while retaining journal-specific wording.")
        self.assertEqual((self.project / "08_submission/integration/full_manuscript.md").read_text(), merged)
        self.assertTrue(journalworkspace.verify(self.project)[0])

    def test_failed_docx_rebuild_preserves_existing_file(self):
        from manuscript.build_docx import build
        target = self.text("08_submission/bundle/preserved.docx", "User's previous output bytes")
        source = self.text("08_submission/integration/title_page.md", "A valid source")
        args = argparse.Namespace(input=[source], output=target, kind="title_page", style_config=None,
                                  bibliography=None, csl=None)
        with patch("manuscript.build_docx.project_root", return_value=self.project), \
             patch("manuscript.build_docx.load_style", return_value={"all_text_black": True, "external_hyperlinks": False}), \
             patch("manuscript.build_docx.run_pandoc", side_effect=RuntimeError("Synthetic compiler failure")), \
             contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(build(args), 2)
        self.assertEqual(target.read_text(), "User's previous output bytes")


if __name__ == "__main__":
    unittest.main(verbosity=2)
