#!/usr/bin/env python3
"""Build only a selected-journal-required checklist from its real item/location source."""
import argparse
import json
import re
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from docx import Document
from docx.shared import Inches
from docx.oxml import OxmlElement
from wfcore import paths
from wfcore.revision import guard_outputs, record_build
from manuscript.build_docx import load_style, normalize_docx, audit_docx


def build(project, source, output, style_path):
    guard_outputs(project, [output])
    data = json.loads(source.read_text(encoding="utf-8"))
    if not data.get("required_by") or not data.get("source") or not data.get("rows"): raise ValueError("only a genuinely journal-required, sourced checklist with actual item locations may be built")
    headers = data.get("columns", [])
    if not headers or any(len(row) != len(headers) for row in data["rows"]): raise ValueError("checklist rows must match declared columns")
    if re.search(r"(?:[A-Za-z]:[/\\]|\.wf/|project/|tools/)", json.dumps(data["rows"])): raise ValueError("internal paths do not belong in a submission checklist; cite page/section/table locations")
    style = load_style(style_path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="medpaper-checklist-", dir=output.parent) as folder:
        candidate = Path(folder)/output.name
        doc = Document(); doc.add_heading(data.get("title", "Reporting checklist"), 0)
        table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
        width = (8.2677-2*float(style["margins_in"]))
        proportions = data.get("widths", [1]*len(headers))
        if len(proportions) != len(headers) or any(value <= 0 for value in proportions): raise ValueError("invalid checklist column weights")
        for index, (column, weight) in enumerate(zip(table.columns, proportions)):
            column.width = Inches(width*weight/sum(proportions)); table.rows[0].cells[index].width = column.width
            table.rows[0].cells[index].text = str(headers[index])
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for values in data["rows"]:
            cells = table.add_row().cells
            for index, value in enumerate(values): cells[index].text = str(value); cells[index].width = table.columns[index].width
        doc.save(candidate); normalize_docx(candidate, style, "checklist_form")
        problems = audit_docx(candidate, style, "checklist_form")
        if problems: raise ValueError("; ".join(problems))
        candidate.replace(output)
    record_build(project, output, [source, style_path, Path(__file__)], entrypoints=[source], tool="tools/manuscript/checklist_docx.py")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True); parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--style-config", type=Path, required=True); parser.add_argument("--project", type=Path)
    args = parser.parse_args()
    try: build((args.project or paths.project_dir()).resolve(), args.input, args.output, args.style_config); print("required checklist Word built; render and inspect before delivery"); return 0
    except (OSError, ValueError, KeyError) as exc: print(str(exc)); return 2


if __name__ == "__main__": raise SystemExit(main())
