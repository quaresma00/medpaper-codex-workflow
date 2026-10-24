#!/usr/bin/env python3
"""Build and audit journal-ready DOCX files with deterministic medical-manuscript styling."""
from __future__ import annotations

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import tomllib
import zipfile
from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE
    from docx.shared import Inches, Pt, RGBColor
except ImportError as exc:  # pragma: no cover - clear deployment failure
    raise SystemExit("python-docx is required; run bootstrap.py to install pinned dependencies") from exc


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools"))
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"w": W, "r": R, "rel": PKG_REL}
TEXT_ROLES = {"manuscript", "title_page", "cover_letter", "supplementary",
              "statements", "figure_legends"}
FORBIDDEN_PARAGRAPH_CONTROLS = (
    "outlineLvl", "keepNext", "keepLines", "pageBreakBefore",
)
THEMATIC_BREAK_RE = re.compile(
    r"(?m)^\s{0,3}(?:(?:\*\s*){3,}|(?:-\s*){3,}|(?:_\s*){3,})\s*$"
)
PIPE_TABLE_DIVIDER_RE = re.compile(
    r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$"
)


def project_root() -> Path:
    raw = os.environ.get("MEDPAPER_PROJECT")
    return Path(raw).resolve() if raw else ROOT / "project"


def planned_figure_ids(project: Path | None = None, scope: str = "all") -> list[str] | None:
    """Return every planned main/supplementary figure id, or None if the plan is unusable."""
    path = (project or project_root()) / "01_protocol/artifact_plan.json"
    journal = (project or project_root()) / "08_submission/integration/artifact_plan.json"
    if journal.is_file():
        path = journal
    if not path.is_file():
        return None
    try:
        plan = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return None
    groups={"all":("main_figures","supp_figures"),"main":("main_figures",),"supp":("supp_figures",)}
    if scope not in groups:
        raise ValueError("Legend scope must be main, supp or all")
    return [str(item.get("id", "")).strip()
            for group in groups[scope]
            for item in plan.get(group, []) if str(item.get("id", "")).strip()]


def figure_legend_cap() -> int | None:
    from wfcore.qualitylimits import legend_limit
    return legend_limit(project_root())

def load_style(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    defaults = {
        "font_family": "Times New Roman", "body_font_pt": 12,
        "title_font_pt": 14, "section_heading_font_pt": 12,
        "subsection_heading_font_pt": 12, "line_spacing": 2.0,
        "paragraph_spacing_after_pt": 0, "margins_in": 1.0,
        "paper_size": "A4", "all_text_black": True,
        "external_hyperlinks": False,
    }
    defaults.update(data)
    project = path.resolve().parent.parent if path.resolve().parent.name == "08_submission" else project_root()
    from wfcore.journalspec import load
    defaults["journal_profile"] = load(project)
    defaults["journal_project"] = str(project)
    return defaults


def _ensure_style(doc, name: str, base: str, font: str, size: float,
                  bold: bool = False, italic: bool = False):
    try:
        style = doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles[base]
    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    ppr = style.element.get_or_add_pPr()
    for tag in (*[f"w:{name}" for name in FORBIDDEN_PARAGRAPH_CONTROLS], "w:numPr"):
        node = ppr.find(qn(tag))
        if node is not None:
            ppr.remove(node)
    return style


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def iter_paragraphs(doc):
    yield from doc.paragraphs
    for table in doc.tables:
        yield from _iter_table_paragraphs(table)
    for section in doc.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def _flatten_hyperlinks(doc) -> None:
    for hyperlink in list(doc.element.xpath(".//w:hyperlink")):
        parent = hyperlink.getparent()
        idx = parent.index(hyperlink)
        for child in list(hyperlink):
            hyperlink.remove(child)
            parent.insert(idx, child)
            idx += 1
        parent.remove(hyperlink)
    for part in doc.part.package.parts:
        for rel_id, rel in list(part.rels.items()):
            if rel.reltype == RELATIONSHIP_TYPE.HYPERLINK:
                part.drop_rel(rel_id)


def _replace_manual_line_breaks(doc) -> None:
    """Convert manual line breaks to real paragraphs without merging address/list lines."""
    for part in doc.part.package.parts:
        root = getattr(part, "element", None)
        if root is None:
            continue
        for element in list(root.iter(qn("w:br"))) + list(root.iter(qn("w:cr"))):
            if element.tag == qn("w:br") and element.get(qn("w:type")) not in (None, "textWrapping"):
                continue
            paragraph = next((node for node in element.iterancestors() if node.tag == qn("w:p")), None)
            if paragraph is None:
                raise ValueError("manual line break outside a paragraph cannot be converted safely")
            current, tail = element, None
            while current is not paragraph:
                parent = current.getparent()
                following = list(parent)[parent.index(current) + 1:]
                clone = deepcopy(parent)
                for child in list(clone):
                    if child.tag not in {qn("w:pPr"), qn("w:rPr")}:
                        clone.remove(child)
                if tail is not None:
                    clone.append(tail)
                for child in following:
                    clone.append(child)
                current, tail = parent, clone
            paragraph.addnext(tail)
            # A section break belongs to the final resulting paragraph only.
            ppr = paragraph.find(qn("w:pPr"))
            sect = None if ppr is None else ppr.find(qn("w:sectPr"))
            if sect is not None:
                ppr.remove(sect)
            element.getparent().remove(element)


def _remove_heading_metadata(paragraph, *, strip_numbering: bool = True) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    tags = [f"w:{name}" for name in FORBIDDEN_PARAGRAPH_CONTROLS]
    if strip_numbering:
        tags.append("w:numPr")
    for tag in tags:
        node = ppr.find(qn(tag))
        if node is not None:
            ppr.remove(node)


def _source_problems(kind: str, inputs: list[Path]) -> list[str]:
    text = "\n\n".join(path.read_text(encoding="utf-8", errors="replace") for path in inputs)
    problems: list[str] = []
    if kind in {"manuscript", "figure_legends"}:
        count = len(re.findall(r"(?mi)^#\s+Figure legends\s*$", text))
        if count != 1:
            problems.append(f"source must contain exactly one '# Figure legends' heading, found {count}")
    if kind == "supplementary":
        lines = text.splitlines()
        if (any(PIPE_TABLE_DIVIDER_RE.match(line) for line in lines) or
                re.search(r"<table\b", text, re.I | re.S) or
                re.search(r"^\s*\+[=-]{3,}(?:\+[=-]{3,})+\+\s*$", text, re.M)):
            problems.append(
                "supplementary Methods contains a table; move it to the planned supplementary "
                "tables workbook and build it with tools/tables/threeline.py"
            )
        if THEMATIC_BREAK_RE.search(text):
            problems.append("supplementary Methods contains a Markdown thematic break/horizontal rule")
    return problems


def _set_run(run, font: str, size: float) -> None:
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    rstyle = rpr.find(qn("w:rStyle"))
    if rstyle is not None and rstyle.get(qn("w:val"), "").casefold() == "hyperlink":
        rpr.remove(rstyle)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)
    for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rfonts.attrib.pop(qn(f"w:{attr}"), None)
    color = rpr.find(qn("w:color"))
    if color is not None:
        color.set(qn("w:val"), "000000")
        for attr in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn(f"w:{attr}"), None)


def normalize_docx(path: Path, style: dict, kind: str) -> None:
    doc = Document(path)
    # A Word save must not silently recompress scientific figure media.
    setting = doc.settings.element.find(qn("w:doNotAutoCompressPictures"))
    if setting is None:
        setting = OxmlElement("w:doNotAutoCompressPictures")
        doc.settings.element.append(setting)
    setting.set(qn("w:val"), "1")
    font = str(style["font_family"])
    body_pt = float(style["body_font_pt"])
    title_pt = float(style["title_font_pt"])
    h1_pt = float(style["section_heading_font_pt"])
    h2_pt = float(style["subsection_heading_font_pt"])
    spacing = float(style["line_spacing"])
    after = float(style["paragraph_spacing_after_pt"])

    body_style = _ensure_style(doc, "Manuscript Body", "Normal", font, body_pt)
    title_style = _ensure_style(doc, "Manuscript Title", "Normal", font, title_pt, bold=True)
    # Both Markdown section levels are flattened onto one Normal-based paragraph
    # style.  Typography may still reflect the journal's H1/H2 sizes, but Word has
    # no outline hierarchy to display or collapse.
    section_style = _ensure_style(doc, "SectionHeading", "Normal", font, h1_pt, bold=True)
    for builtin in ("Normal", "Body Text", "Title", "Subtitle", "Heading 1", "Heading 2",
                    "Heading 3", "Bibliography", "Hyperlink"):
        try:
            s = doc.styles[builtin]
        except KeyError:
            continue
        s.font.name = font
        s.font.color.rgb = RGBColor(0, 0, 0)
        if builtin == "Hyperlink":
            s.font.underline = False
        ppr = s.element.get_or_add_pPr()
        for tag in ("w:outlineLvl", "w:keepNext", "w:keepLines", "w:pageBreakBefore"):
            node = ppr.find(qn(tag))
            if node is not None:
                ppr.remove(node)

    _flatten_hyperlinks(doc)
    _replace_manual_line_breaks(doc)
    from wfcore.journalspec import apply_docx
    apply_docx(doc, style.get("journal_profile", {}), kind)
    _format_reference_spacing(doc, style.get("reference_separator", "space"))
    paras = list(iter_paragraphs(doc))
    nonempty_seen = 0
    for p in paras:
        original = p.style.name if p.style is not None else ""
        text = p.text.strip()
        if text:
            nonempty_seen += 1
        _remove_heading_metadata(p, strip_numbering=False)
        if original == "Title" or (kind == "manuscript" and nonempty_seen == 1):
            p.style = title_style
            size = title_pt
            _remove_heading_metadata(p)
        elif re.match(r"^Heading 1$", original, re.I):
            p.style = section_style
            size = h1_pt
            _remove_heading_metadata(p)
            # Invisible identity, not Word outline hierarchy. Survives normal Word saves.
            bookmark = OxmlElement("w:bookmarkStart")
            bookmark.set(qn("w:id"),str(10000+nonempty_seen))
            slug = re.sub(r"[^a-z0-9]+","_",text.casefold()).strip("_")
            bookmark.set(qn("w:name"),("medpaper_section_"+str(nonempty_seen)+"_"+slug)[:40])
            p._p.insert(1 if p._p.pPr is not None else 0,bookmark)
            end = OxmlElement("w:bookmarkEnd"); end.set(qn("w:id"),str(10000+nonempty_seen));p._p.append(end)
        elif re.match(r"^Heading [2-9]$", original, re.I):
            p.style = section_style
            size = h2_pt
            _remove_heading_metadata(p)
        elif original in {"SectionHeading","Manuscript Title","ManuscriptTitle"}:
            previous_size=next((run.font.size.pt for run in p.runs if run.font.size is not None),h1_pt)
            size = title_pt if "Title" in original else previous_size if previous_size in {h1_pt,h2_pt} else h1_pt
        else:
            if original in {"Normal", "Body Text", "First Paragraph", "Compact"}:
                p.style = body_style
            size = body_pt
        p.paragraph_format.line_spacing = spacing
        p.paragraph_format.space_after = Pt(after)
        for run in p.runs:
            _set_run(run, font, size)

    margin = Inches(float(style["margins_in"]))
    for section_index, section in enumerate(doc.sections):
        landscape = section.orientation == WD_ORIENT.LANDSCAPE
        section.top_margin = section.bottom_margin = margin
        section.left_margin = section.right_margin = margin
        if str(style["paper_size"]).upper() == "LETTER":
            section.page_width, section.page_height = Inches(8.5), Inches(11)
        else:
            section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
        if landscape:
            section.page_width, section.page_height = section.page_height, section.page_width
        section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
        if style.get("page_numbers", True):
            section.footer.is_linked_to_previous = section_index > 0
            has_page = any("PAGE" in (node.text or "").split() for node in section.footer._element.iter(qn("w:instrText"))) or any(
                "PAGE" in node.get(qn("w:instr"), "").split() for node in section.footer._element.iter(qn("w:fldSimple")))
            if not has_page:
                paragraph = section.footer.paragraphs[0]
                field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE")
                paragraph._p.append(field)
            numbering = section._sectPr.find(qn("w:pgNumType"))
            if numbering is not None:
                numbering.attrib.pop(qn("w:start"), None)

    doc.save(path)
    _sanitize_package_xml(path, font)
    from wfcore.determinism import normalize
    normalize(path)


def _format_reference_spacing(doc, separator: str) -> None:
    """Preserve reference identity/formatting; normalize only numeric lead spacing."""
    if separator not in {"space", "tab"}:
        raise ValueError("reference_separator must be space or tab, based on actual journal rules")
    for paragraph in doc.paragraphs:
        if paragraph.style is None or paragraph.style.name != "Bibliography":
            continue
        match = re.match(r"^\s*(\[\d+\]|\d+[.)]?)[ \t\u00a0]*", paragraph.text)
        if not match or not paragraph.runs:
            continue
        label, remaining = match.group(1), match.end()
        for run in paragraph.runs:
            length = len(run.text)
            if remaining:
                run.text = run.text[min(length, remaining):]
                remaining = max(0, remaining - length)
        first = paragraph.runs[0]
        first.text = label + ("\t" if separator == "tab" else " ") + first.text
        paragraph.paragraph_format.left_indent = Inches(0.5 if separator == "tab" else 0)
        paragraph.paragraph_format.first_line_indent = Inches(-0.5 if separator == "tab" else 0)
        tabs = paragraph._p.get_or_add_pPr().find(qn("w:tabs"))
        if tabs is not None:
            tabs.getparent().remove(tabs)
        if separator == "tab":
            paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(0.5))


def _sanitize_package_xml(path: Path, font: str) -> None:
    """Remove theme residue and every Word paragraph control the user prohibited."""
    with zipfile.ZipFile(path, "r") as src:
        entries = [(item, src.read(item.filename)) for item in src.infolist()]
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for item, blob in entries:
            if item.filename.startswith("word/") and item.filename.endswith(".xml"):
                # Word's mc:Ignorable attributes contain namespace-prefix strings; parsing
                # and reserializing them with ElementTree can rename prefixes and make Word
                # report a corrupt file. Byte replacement preserves the package exactly.
                blob = re.sub(br"Aptos(?: Display)?", font.encode("utf-8"), blob,
                              flags=re.I)
                for tag in (*FORBIDDEN_PARAGRAPH_CONTROLS, "pBdr"):
                    name = tag.encode("ascii")
                    blob = re.sub(rb"<w:" + name + rb"\b[^>]*/>", b"", blob)
                    blob = re.sub(rb"<w:" + name + rb"\b[^>]*>.*?</w:" + name + rb">",
                                  b"", blob, flags=re.S)
            out.writestr(item, blob)
    tmp.replace(path)


def docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    return "\n".join("".join(t.text or "" for t in p.findall(".//w:t", NS))
                       for p in root.findall(".//w:p", NS))


def _canon_figure_id(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().rstrip(".")).casefold()


def _legend_problems(root, expected: list[str], cap: int) -> list[str]:
    paragraphs: list[tuple[str, str]] = []
    for paragraph in root.findall(".//w:p", NS):
        paragraph_text = "".join(t.text or "" for t in paragraph.findall(".//w:t", NS)).strip()
        pstyle = paragraph.find("w:pPr/w:pStyle", NS)
        style_id = "" if pstyle is None else pstyle.get(f"{{{W}}}val", "")
        if paragraph_text:
            paragraphs.append((paragraph_text, style_id))
    try:
        start = next(i for i, (value, paragraph_style) in enumerate(paragraphs)
                     if value.casefold() == "figure legends" and paragraph_style=="SectionHeading") + 1
    except StopIteration:
        return ["Figure legends heading is missing"]
    headings: list[tuple[int, str, str]] = []
    for i in range(start, len(paragraphs)):
        value, style_id = paragraphs[i]
        match = re.match(r"^(Figure\s+S?\d+)\.?\s*(.*)$", value, re.I)
        if match and style_id == "SectionHeading":
            headings.append((i, match.group(1), match.group(2)))
    blocks: dict[str, tuple[int, int]] = {}
    for j, (idx, figure_id, suffix) in enumerate(headings):
        end = headings[j + 1][0] if j + 1 < len(headings) else len(paragraphs)
        content = "\n".join([suffix] + [value for value, _ in paragraphs[idx + 1:end]]).strip()
        blocks[_canon_figure_id(figure_id)] = (
            len(content), len(re.findall(r"\b[A-Za-z][A-Za-z0-9'-]*\b", content)))
    wanted = {_canon_figure_id(value): value for value in expected}
    problems: list[str] = []
    seen = [_canon_figure_id(item[1]) for item in headings]
    duplicates = sorted({item for item in seen if seen.count(item) > 1})
    if duplicates:
        problems.append("duplicate figure legend heading(s): " + ", ".join(duplicates))
    for idx, figure_id, _ in headings:
        if idx + 1 < len(paragraphs):
            first_body = paragraphs[idx + 1][0]
            if re.match(rf"^{re.escape(figure_id)}\b", first_body, re.I):
                problems.append(f"{figure_id} is repeated at the start of its legend body")
    for key, display in wanted.items():
        if key not in blocks:
            problems.append(f"Figure legends lacks {display}")
            continue
        chars, words = blocks[key]
        if chars < 40:
            problems.append(f"{display} legend is only {chars} characters")
        if cap is not None and words > cap:
            problems.append(f"{display} legend is {words} words (maximum {cap})")
    extras = sorted(set(blocks) - set(wanted))
    if extras:
        problems.append("unplanned figure legend(s): " + ", ".join(extras))
    return problems


def audit_docx(path: Path, style: dict, kind: str = "",
               planned_figures: list[str] | None = None,
               legend_cap: int | None = None) -> list[str]:
    problems: list[str] = []
    if not path.is_file():
        return [f"missing DOCX: {path}"]
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip(): return [f"{path.name}: corrupt Office container"]
    except (OSError, zipfile.BadZipFile) as exc:
        return [f"{path.name}: invalid DOCX ({exc})"]
    from wfcore.journalspec import docx_problems
    problems += docx_problems(path, style.get("journal_profile", {}), Path(style.get("journal_project", project_root())), kind)
    from wfcore.wordtables import table_problems
    problems.extend(table_problems(Document(path)))
    for paragraph in Document(path).paragraphs:
        if paragraph.style is not None and paragraph.style.name == "Bibliography":
            match = re.match(r"^\s*(?:\[\d+\]|\d+[.)]?)([ \t\u00a0]*)", paragraph.text)
            expected = "\t" if style.get("reference_separator", "space") == "tab" else " "
            if match and match.group(1) != expected:
                problems.append(f"{path.name}: actual numeric reference lead must use exactly one {style.get('reference_separator', 'space')}")
    try:
        with zipfile.ZipFile(path) as zf:
            bad = zf.testzip()
            if bad:
                return [f"{path.name}: corrupt at {bad}"]
            xmls = {n: zf.read(n) for n in zf.namelist()
                    if n.startswith("word/") and n.endswith((".xml", ".rels"))}
    except zipfile.BadZipFile:
        return [f"{path.name}: invalid ZIP/DOCX"]

    document = xmls.get("word/document.xml", b"")
    try:
        root = ET.fromstring(document)
    except ET.ParseError as exc:
        return [f"{path.name}: invalid document.xml ({exc})"]
    if root.findall(".//w:hyperlink", NS):
        problems.append(f"{path.name}: hyperlink elements remain")
    if any("HYPERLINK" in (node.text or "").upper()
           for node in root.findall(".//w:instrText", NS)):
        problems.append(f"{path.name}: hyperlink field code remains")
    for name, blob in xmls.items():
        if name.endswith(".rels"):
            try:
                relroot = ET.fromstring(blob)
            except ET.ParseError:
                continue
            for rel in relroot.findall("rel:Relationship", NS):
                if rel.get("TargetMode") == "External" and "hyperlink" in rel.get("Type", "").lower():
                    problems.append(f"{path.name}: external hyperlink relationship remains")
    for name, blob in xmls.items():
        if not name.endswith(".xml"):
            continue
        part = ET.fromstring(blob)
        if any(node.get(f"{{{W}}}type") in (None, "textWrapping")
               for node in part.findall(".//w:br", NS)) or part.findall(".//w:cr", NS):
            problems.append(f"{path.name}: manual line-break control remains in {name}")
        if any("\u2193" in (node.text or "") for node in part.findall(".//w:t", NS)):
            problems.append(f"{path.name}: forbidden down-arrow character remains in {name}")
        for tag in FORBIDDEN_PARAGRAPH_CONTROLS:
            if part.findall(f".//w:{tag}", NS):
                problems.append(f"{path.name}: forbidden {tag} paragraph control remains in {name}")
        if part.findall(".//w:pBdr", NS):
            problems.append(f"{path.name}: paragraph border/horizontal-rule residue remains in {name}")
    expected_font = str(style["font_family"])
    if expected_font.casefold() == "times new roman" and any(b"Aptos" in b for b in xmls.values()):
        problems.append(f"{path.name}: Aptos remains in Word XML")
    for p in root.findall(".//w:p", NS):
        ppr = p.find("w:pPr", NS)
        pstyle = None if ppr is None else ppr.find("w:pStyle", NS)
        style_id = "" if pstyle is None else pstyle.get(f"{{{W}}}val", "")
        allowed_pts = ({float(style["title_font_pt"])} if style_id == "ManuscriptTitle" else
                       {float(style["section_heading_font_pt"]),
                        float(style["subsection_heading_font_pt"])} if style_id == "SectionHeading" else
                       {float(style["body_font_pt"])})
        paragraph_text = "".join(t.text or "" for t in p.findall(".//w:t", NS)).strip()
        spacing_node = None if ppr is None else ppr.find("w:spacing", NS)
        if paragraph_text and (spacing_node is None or
                spacing_node.get(f"{{{W}}}line") !=
                str(int(round(float(style["line_spacing"]) * 240))) or
                spacing_node.get(f"{{{W}}}after", "0") !=
                str(int(round(float(style["paragraph_spacing_after_pt"]) * 20)))):
            problems.append(f"{path.name}: inconsistent line/paragraph spacing")
        for run in p.findall(".//w:r", NS):
            if not "".join(t.text or "" for t in run.findall(".//w:t", NS)).strip():
                continue
            rpr = run.find("w:rPr", NS)
            rf = None if rpr is None else rpr.find("w:rFonts", NS)
            color = None if rpr is None else rpr.find("w:color", NS)
            size = None if rpr is None else rpr.find("w:sz", NS)
            if rf is None or any(rf.get(f"{{{W}}}{a}") != expected_font for a in ("ascii", "hAnsi")):
                problems.append(f"{path.name}: a text run does not explicitly use {expected_font}")
                break
            if color is None or color.get(f"{{{W}}}val", "").upper() not in {"000000", "AUTO"}:
                problems.append(f"{path.name}: a text run is not explicitly black")
                break
            if size is None or size.get(f"{{{W}}}val", "") not in {
                    str(int(round(value * 2))) for value in allowed_pts}:
                problems.append(f"{path.name}: inconsistent size in {style_id or 'body'} paragraph")
                break
        if style_id.casefold().startswith("heading"):
            problems.append(f"{path.name}: built-in outline heading style remains on visible text")
            break
    text = docx_text(path)
    separator=style.get("reference_number_separator")
    if separator is not None and kind=="manuscript":
        try:
            from wfcore.submission import docx_sections
            _,sections,_=docx_sections(path)
            references=sections.get("References","")
            observed=re.findall(r"(?m)^\s*\d+([.)\]])\s+",references)
            if not observed:
                # Word stores automatic list labels in numbering.xml, not in paragraph text.
                with zipfile.ZipFile(path) as archive:
                    numbering=ET.fromstring(archive.read("word/numbering.xml")) if "word/numbering.xml" in archive.namelist() else None
                    styles=ET.fromstring(archive.read("word/styles.xml"))
                nums={n.get(f"{{{W}}}numId"):n.find("w:abstractNumId",NS).get(f"{{{W}}}val") for n in numbering.findall("w:num",NS)} if numbering is not None else {}
                abstracts={n.get(f"{{{W}}}abstractNumId"):n for n in numbering.findall("w:abstractNum",NS)} if numbering is not None else {}
                active=False
                for paragraph in root.findall(".//w:p",NS):
                    paragraph_text="".join(t.text or "" for t in paragraph.findall(".//w:t",NS)).strip()
                    ps=paragraph.find("w:pPr/w:pStyle",NS)
                    sid=ps.get(f"{{{W}}}val","") if ps is not None else ""
                    if sid in {"SectionHeading","Heading1"} and paragraph_text.casefold()=="references":active=True;continue
                    if active and sid in {"SectionHeading","Heading1"}:break
                    if not active:continue
                    numpr=paragraph.find("w:pPr/w:numPr",NS)
                    if numpr is None:
                        st=styles.find(f"w:style[@w:styleId='{sid}']",NS)
                        numpr=st.find("w:pPr/w:numPr",NS) if st is not None else None
                    if numpr is None:continue
                    numid=numpr.find("w:numId",NS);level=numpr.find("w:ilvl",NS)
                    abstract=abstracts.get(nums.get(numid.get(f"{{{W}}}val") if numid is not None else ""))
                    if abstract is None:continue
                    lvl=abstract.find(f"w:lvl[@w:ilvl='{level.get(f'{{{W}}}val') if level is not None else '0'}']",NS)
                    marker=lvl.find("w:lvlText",NS) if lvl is not None else None
                    match=re.search(r"%\d([.)\]])",marker.get(f"{{{W}}}val","") if marker is not None else "")
                    if match:observed.append(match[1])
            if not observed:
                problems.append(f"{path.name}: reference numbering separator not observable in actual text; inspect Word numbering/CSL")
            elif any(value!=separator for value in observed):
                problems.append(f"{path.name}: actual reference numbering separator differs from journal rule {separator!r}")
        except (OSError,ValueError,KeyError) as exc:
            problems.append(f"{path.name}: reference delimiter inspection failed: {exc}")
    if "\u2193" in text:
        problems.append(f"{path.name}: forbidden down-arrow character remains")
    if kind == "manuscript":
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        exact = {name: [i for i, line in enumerate(lines) if line.casefold() == name]
                 for name in ("abstract", "references", "figure legends")}
        keywords = [i for i, line in enumerate(lines) if line.casefold().startswith("keywords:")]
        for name, indices in exact.items():
            if len(indices) != 1:
                problems.append(f"{path.name}: expected one {name.title()} heading, found {len(indices)}")
        if len(keywords) != 1:
            problems.append(f"{path.name}: expected one Keywords line, found {len(keywords)}")
        order = [exact[name][0] if len(exact[name]) == 1 else -1
                 for name in ("abstract",)] + [keywords[0] if len(keywords) == 1 else -1] + [
                 exact[name][0] if len(exact[name]) == 1 else -1
                 for name in ("references", "figure legends")]
        if all(x >= 0 for x in order) and order != sorted(order):
            problems.append(f"{path.name}: Abstract/Keywords/References/Figure legends order is wrong")
        if planned_figures is None:
            problems.append(f"{path.name}: artifact_plan.json is missing or invalid")
        else:
            problems.extend(f"{path.name}: {problem}" for problem in
                            _legend_problems(root, planned_figures, legend_cap))
    elif kind == "figure_legends":
        count = sum(1 for line in docx_text(path).splitlines()
                    if line.strip().casefold() == "figure legends")
        if count != 1:
            problems.append(f"{path.name}: expected one Figure legends heading, found {count}")
    if kind == "supplementary" and root.findall(".//w:tbl", NS):
        problems.append(
            f"{path.name}: supplementary Methods contains a table; package it as a separate "
            "three-line supplementary table"
        )
    return sorted(set(problems))


def run_pandoc(inputs: list[Path], output: Path, bibliography: Path | None,
               csl: Path | None) -> None:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("pandoc is required to render Markdown and citations")
    cmd = [pandoc, *map(str, inputs)]
    if bibliography:
        cmd += ["--citeproc", f"--bibliography={bibliography}"]
    if csl:
        cmd += [f"--csl={csl}"]
    cmd += ["--metadata=link-citations:false", "-o", str(output)]
    from wfcore.execution import run_logged
    proc = run_logged(cmd,project=project_root(),timeout=180)
    combined = (proc.stdout or "") + (proc.stderr or "")
    if proc.returncode:
        raise RuntimeError((combined.strip() or f"pandoc exited {proc.returncode}")[-2000:]+f"; full log: {proc.log_path}")
    if re.search(r"could not find citation|unresolved", combined, re.I):
        raise RuntimeError("pandoc reported an unresolved citation: " + combined.strip())


def build(args) -> int:
    from wfcore.dependencies import guard_build
    from wfcore.revision import guard_outputs, record_build
    guard_build(project_root())
    guard_outputs(project_root(), [args.output])
    style = load_style(args.style_config)
    if style.get("all_text_black") is not True or style.get("external_hyperlinks") is not False:
        print("style config must require black text and prohibit external hyperlinks", file=sys.stderr)
        return 2
    args.output.parent.mkdir(parents=True, exist_ok=True)
    from wfcore.packagefreeze import _sha256
    tables_path = getattr(args, "tables_json", None)
    tables, table_inputs = [], []
    if tables_path:
        from wfcore.wordtables import load_spec
        tables, table_inputs = load_spec(project_root(), tables_path)
    profile_path = project_root()/"08_submission/target_journal.json"
    profile_inputs = [profile_path] if profile_path.is_file() else []
    author_input = project_root()/"00_input/author_info.json"
    if style.get("journal_profile", {}).get("blinded_main_manuscript") and author_input.is_file():
        profile_inputs.append(author_input)
    patch_path = getattr(args, "csl_patch", None) or project_root()/"08_submission/integration/csl_patch.json"
    patch_path = patch_path if patch_path.is_file() else None
    if patch_path:
        if not args.csl: raise ValueError("CSL patch needs an explicitly sourced base --csl")
        profile_inputs.append(patch_path)
    actual_inputs = [p for p in [*args.input,args.style_config,args.bibliography,args.csl,*table_inputs,*profile_inputs] if p is not None]
    before = {p:_sha256(p) for p in actual_inputs}
    try:
        source_problems = _source_problems(args.kind, args.input)
        if source_problems:
            raise ValueError("; ".join(source_problems))
        # Validate a temporary candidate; a failed rebuild must preserve the user's file.
        with tempfile.TemporaryDirectory(prefix="medpaper-docx-", dir=args.output.parent) as folder:
            candidate = Path(folder) / args.output.name
            used_csl = args.csl
            if patch_path:
                from wfcore.cslpatch import apply
                used_csl = Path(folder)/"journal.csl"
                apply(args.csl, patch_path, used_csl)
            run_pandoc(args.input, candidate, args.bibliography, used_csl)
            if tables:
                from wfcore.wordtables import append_tables
                document = Document(candidate)
                append_tables(document, project_root(), tables, style)
                document.save(candidate)
            normalize_docx(candidate, style, args.kind)
            problems = audit_docx(candidate, style, args.kind,
                                  planned_figure_ids(scope=getattr(args,"legend_scope",None) or style.get("figure_legend_scope","all"))
                                  if args.kind in {"manuscript","figure_legends"} else None,
                                  figure_legend_cap())
            if problems:
                raise ValueError("DOCX audit failed: " + "; ".join(problems))
            if any(not p.is_file() or _sha256(p)!=h for p,h in before.items()):
                raise ValueError("DOCX input changed during build; previous output preserved")
            candidate.replace(args.output)
    except Exception as exc:  # noqa: BLE001
        print(f"DOCX build failed: {exc}", file=sys.stderr)
        return 2
    project = project_root()
    plan = project/"08_submission/integration/artifact_plan.json"
    if not plan.is_file():
        plan = project/"01_protocol/artifact_plan.json"
    observed = [*args.input,args.style_config,args.bibliography,args.csl,*table_inputs,*profile_inputs,Path(__file__),
                ROOT/"tools/wfcore/journalspec.py",ROOT/"tools/wfcore/determinism.py",
                ROOT/"tools/wfcore/qualitylimits.py",
                ROOT/"tools/wfcore/cslpatch.py",
                ROOT/"tools/wfcore/wordtables.py",
                project.parent/"pipeline/pipeline.toml", plan if plan.is_file() else None]
    observed = [p for p in observed if p is not None and p.is_file()]
    record_build(project, args.output, observed,
                 entrypoints=[*args.input, *([tables_path] if tables_path else [])],tool="tools/manuscript/build_docx.py",
                 parameters={"kind":args.kind,"style":str(args.style_config),"bibliography":str(args.bibliography or ""),
                             "csl":str(args.csl or ""),"tables_json":str(tables_path or ""),
                             "legend_scope":getattr(args,"legend_scope",None) or style.get("figure_legend_scope","all")})
    print(f"built and audited {args.kind} -> {args.output}")
    return 0


def finalize(args) -> int:
    """Clean task-owned Word after layout repair; never recapture changed text."""
    from wfcore.revision import guard_outputs
    from wfcore.buildproof import content_digest
    project = project_root();guard_outputs(project,[args.output])
    before = content_digest(args.output)
    style = load_style(args.style_config)
    with tempfile.TemporaryDirectory(prefix="medpaper-finalize-",dir=args.output.parent) as folder:
        candidate=Path(folder)/args.output.name;shutil.copy2(args.output,candidate)
        normalize_docx(candidate,style,args.kind)
        if content_digest(candidate) != before:
            raise ValueError("Final cleanup changed visible content; prior Word preserved")
        problems=audit_docx(candidate,style,args.kind,planned_figure_ids() if args.kind=="manuscript" else None,figure_legend_cap())
        if problems:
            raise ValueError("Final cleanup QA failed: "+"; ".join(problems))
        candidate.replace(args.output)
    print("Word cleanup complete; render this final file, then bind baseline/portal and user approval")
    return 0


def _resolve_manifest_file(manifest: Path, raw: str) -> Path:
    p = Path(raw)
    if p.is_absolute():
        return p
    project = project_root()
    return project / p


def audit_manifest(args) -> int:
    style = load_style(args.style_config)
    manifest = json.loads(args.manifest.read_text(encoding="utf-8"))
    problems: list[str] = []
    roles: dict[str, list[Path]] = {}
    planned = planned_figure_ids()
    legend_cap = figure_legend_cap()
    for item in manifest.get("items", []):
        role = str(item.get("role", "")).strip()
        path = _resolve_manifest_file(args.manifest, str(item.get("file", "")))
        roles.setdefault(role, []).append(path)
        if role in TEXT_ROLES and path.suffix.lower() != ".docx":
            problems.append(f"{role}: narrative upload must be DOCX, got {path.name}")
        if path.suffix.lower() == ".docx":
            scope=item.get("legend_scope",style.get("figure_legend_scope","all"))
            wanted=planned_figure_ids(scope=scope) if role in {"manuscript","figure_legends"} else planned
            problems += audit_docx(path, style, role, wanted, legend_cap)
    for role in ("manuscript", "title_page", "cover_letter"):
        if not any(p.suffix.lower() == ".docx" for p in roles.get(role, [])):
            problems.append(f"manifest lacks a DOCX item for role '{role}'")
    if (project_root() / "08_submission/integration/supplementary_methods.md").exists():
        if not any(p.suffix.lower() == ".docx" for p in roles.get("supplementary", [])):
            problems.append("supplementary_methods.md exists but no supplementary DOCX is listed")
    if problems:
        print("submission DOCX audit failed:\n  " + "\n  ".join(sorted(set(problems))))
        return 2
    print(f"submission DOCX audit passed: {sum(len(v) for v in roles.values())} manifest item(s)")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser(description="build or audit deterministic journal DOCX files")
    sub = ap.add_subparsers(dest="command", required=True)
    p = sub.add_parser("build")
    p.add_argument("--kind", required=True,
                   choices=["manuscript", "title_page", "cover_letter", "supplementary",
                            "statements", "checklist_form", "coi_forms", "figure_legends", "tables"])
    p.add_argument("--input", type=Path, action="append", required=True)
    p.add_argument("--output", type=Path, required=True)
    p.add_argument("--style-config", type=Path, required=True)
    p.add_argument("--bibliography", type=Path)
    p.add_argument("--csl", type=Path)
    p.add_argument("--csl-patch", type=Path, help="local formatting-only patch; never edit a shared CSL or reference metadata")
    p.add_argument("--tables-json", type=Path, help="declared XLSX sheets/ranges/orientations appended as native three-line tables")
    p.add_argument("--legend-scope",choices=["main","supp","all"],help="actual journal legend inventory")
    p.set_defaults(fn=build)
    p = sub.add_parser("finalize",help="clean after Word layout fixes without changing content")
    p.add_argument("--kind",required=True)
    p.add_argument("--output",type=Path,required=True)
    p.add_argument("--style-config",type=Path,required=True)
    p.set_defaults(fn=finalize)
    p = sub.add_parser("audit")
    p.add_argument("--manifest", type=Path, required=True)
    p.add_argument("--style-config", type=Path, required=True)
    p.set_defaults(fn=audit_manifest)
    args = ap.parse_args()
    return args.fn(args)


if __name__ == "__main__":
    raise SystemExit(main())
