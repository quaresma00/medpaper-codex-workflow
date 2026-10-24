"""Synthetic regressions for the user-maintained Kiro 1.9 concepts adapted to Codex."""
import json
import os
from pathlib import Path
import shutil
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from wfcore import registry, revision, reviewrounds, scientificfreeze, presentationdelta as delta
from wfcore.readiness import atomic_json
from wfcore.state import State
from wfcore.packagefreeze import _sha256, write_freeze, verify_freeze
from wfcore.determinism import normalize, content_sha256
from wfcore import journalspec, journalworkspace, finalcopy
from wfcore.checks import Ctx, get, load_all
from manuscript.build_docx import normalize_docx, audit_docx, load_style
from test_source_aligned import DEFAULT_STYLE


class NewMigrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="medpaper-v113-"); self.addCleanup(self.temp.cleanup)
        self.project = Path(self.temp.name); self.pipeline = registry.load(); load_all()
        self.state = State(self.project).create("medpaper", "1.12.0", "S19_human_review")
        self.data("01_protocol/artifact_plan.json", {"main_figures": [], "main_tables": []})
        for rel in ("07_manuscript/full_manuscript.md", "07_manuscript/methods.md", "05_figures/legends.md", "04_tables/table_captions.md"):
            self.text(rel, "Colour and centre labels in a synthetic fixture. N = 162.")

    def text(self, rel, value):
        path = self.project/rel; path.parent.mkdir(parents=True, exist_ok=True); path.write_text(value, encoding="utf-8"); return path

    def data(self, rel, value): atomic_json(self.project/rel, value); return self.project/rel

    def freeze(self):
        from test_review_rounds import ReviewRoundTests
        record = reviewrounds.start(self.project, "manuscript")
        headings = ("Reader comprehension", "Medical presentation")
        report = self.text("temp/synthetic-review.md", "# Synthetic lifecycle fixture\n" + "\n".join(
            f"\n## {heading}\n\n" + "This is controlled test evidence for synthetic artifacts, not a real independent scientific audit. "*2 for heading in headings) +
            "\n```json\n"+json.dumps({"reviewed_artifacts":record["reviewed_artifacts"]})+"\n```\n")
        reviewrounds.attach(self.project,"manuscript",report); reviewrounds.close(self.project,"manuscript",[],[])
        reviewrounds.decide(self.project,"manuscript","no_further_review","我确认这一版")
        return scientificfreeze.write(self.project,"1.12.0")[1]

    def round(self):
        baseline = revision.snapshot(self.project)
        return {"schema_version":2,"round_id":"R001","status":"active","review_stage":"S24_package_human_review",
                "presentation_candidate":True,"baseline":baseline,"presentation_baseline":delta.snapshot(self.project,baseline),
                "items":[{"kind":"methods-wording","change_type":"wording","affected_sources":["07_manuscript/methods.md"]}]}

    def test_spelling_and_spacing_preserve_order_and_clinical_tokens(self):
        before = "Colour at the centre: n = 1,234; P < 0.001; ICD-10; 5-minute exposure."
        after = "Color at the center: n = 1234; P < .001; ICD-10; 5-minute exposure."
        self.assertEqual(delta.canonical(before,self.project),delta.canonical(after,self.project))
        self.assertNotEqual(delta.canonical(before,self.project),delta.canonical(after.replace("5-minute","10-minute"),self.project))

    def test_negation_and_same_numbers_cannot_be_presentation(self):
        self.text("07_manuscript/methods.md","Risk increased for 162 people.")
        record = self.round(); self.text("07_manuscript/methods.md","Risk did not increase for 162 people.")
        self.assertEqual(delta.classify(self.project,record)["class"],"scientific")

    def test_unresolved_paraphrase_needs_real_executor_meaning_review(self):
        record=self.round(); self.text("07_manuscript/methods.md","A synthetic cohort comprised 162 individuals at a centre.")
        self.assertEqual(delta.classify(self.project,record)["class"],"needs_executor_meaning_review")

    def test_data_change_cannot_be_classified_as_layout(self):
        self.text("03_analysis/code/a.py","# synthetic original\n")
        record=self.round(); self.text("03_analysis/code/a.py","# changed producer\n")
        self.assertEqual(delta.classify(self.project,record)["class"],"scientific")

    def test_freeze_and_old_round_survive_proved_spelling_amendment(self):
        freeze=self.freeze(); record=self.round()
        self.text("07_manuscript/full_manuscript.md","Color and center labels in a synthetic fixture. N = 162.")
        self.text("07_manuscript/methods.md","Color and center labels in a synthetic fixture. N = 162.")
        self.assertEqual(delta.classify(self.project,record)["class"],"presentation")
        record.update(status="complete",validation={"ok":True},content_classification={"class":"presentation"})
        self.data(".wf/revisions/R001.json",record); delta.commit(self.project,record,"Synthetic spelling amendment, not a new scientific approval")
        self.assertTrue(scientificfreeze.verify(self.project)[0],scientificfreeze.verify(self.project)[1])
        self.assertEqual(reviewrounds.approved(self.project,"manuscript")["id"],"M001")
        self.assertEqual(json.loads((self.project/scientificfreeze.FREEZE_REL).read_text())["freeze_id"],freeze["freeze_id"])

    def test_no_zip_freeze_can_initialise_journal_workspace(self):
        self.freeze(); self.data("08_submission/target_journal.json",{"journal":"Synthetic venue","issn":"0000-0000"})
        path, data, created=journalworkspace.initialise(self.project)
        self.assertTrue(created); self.assertEqual(data["review_package_id"],"M001")
        self.assertTrue(journalworkspace.verify(self.project)[0])

    def test_old_completed_gates_are_pinned_but_current_stage_is_not(self):
        self.state.data["stages"]["S08_methods"]={"status":"done"}; self.state.save()
        self.state.migrate_pipeline("1.13.0",[stage.id for stage in self.pipeline.stages])
        self.assertTrue(self.state.pin_excludes("S08_methods","1.13.0"))
        self.assertFalse(self.state.pin_excludes("S19_human_review","1.13.0"))
        self.state.migrate_pin("1.13.0"); self.assertFalse(self.state.pin_excludes("S08_methods","1.13.0"))

    def test_office_normalization_is_deterministic_and_does_not_erase_authors(self):
        from datetime import datetime, timezone
        first=self.project/"first.docx"; second=self.project/"second.docx"
        doc=Document(); doc.add_paragraph("Synthetic result 162"); doc.core_properties.author="Fixture Author"
        doc.save(first); doc.core_properties.modified=datetime(2020,1,1,tzinfo=timezone.utc); doc.save(second)
        normalize(first); normalize(second)
        self.assertEqual(first.read_bytes(),second.read_bytes()); self.assertEqual(Document(first).core_properties.author,"Fixture Author")
        before=first.read_bytes(); normalize(first); self.assertEqual(first.read_bytes(),before)

    def test_generated_xlsx_is_byte_stable_and_both_binary_levels_are_counted(self):
        from tables.threeline import write_table
        from analysis.results import baseline_counts
        path=self.project/"t.xlsx"
        options=dict(path=path,sheet="Table 1",title="Table 1. Synthetic baseline",header=["Group","Count"],rows=[["Female",1],["Male",2]],footnotes=["Synthetic counts."])
        write_table(**options); before=path.read_bytes(); write_table(**options); self.assertEqual(before,path.read_bytes())
        counts=baseline_counts(["F","M",None,"M"],["A","A","B","B"])
        self.assertEqual(counts["overall"]["counts"],{"F":1,"M":2}); self.assertEqual(counts["overall"]["missing"],1)
        self.assertEqual(counts["groups"][1]["counts"]["F"],0)

    def test_prose_gate_excludes_declarations_but_catches_report_lists(self):
        path="07_manuscript/full_manuscript.md"
        self.text(path,"# Methods\n\n- First method\n- Second method\n\n# Declarations and Statements\n\n- Author A\n- Author B\n")
        ctx=Ctx(self.pipeline,self.state,self.project,self.pipeline.stage("S17_assemble"),{"path":path})
        self.assertFalse(get("prose_not_lists")(ctx).ok)
        self.text(path,"# Methods\n\nThe cohort followed a stated protocol.\n\n# Declarations and Statements\n\n- Author A\n- Author B\n")
        self.assertTrue(get("prose_not_lists")(ctx).ok)

    def test_first_citation_and_panel_coverage(self):
        self.data("01_protocol/artifact_plan.json",{"main_figures":[{"id":"Figure 1","panels":["A","B"]}]})
        self.text("07_manuscript/full_manuscript.md","# Results\n\nFigure 2 precedes Figure 1A.")
        ctx=Ctx(self.pipeline,self.state,self.project,self.pipeline.stage("S17_assemble"),{})
        self.assertFalse(get("display_numbering_consistent")(ctx).ok)
        self.text("07_manuscript/full_manuscript.md","# Results\n\nFigure 1 shows the result, followed by Figure 2.")
        self.assertTrue(get("display_numbering_consistent")(ctx).ok)

    def test_profile_maps_word_heading_keywords_and_counts_actual_sections(self):
        target={"main_text_headings":{"Methods":"Patients and methods"},"keywords":{"min":3,"max":6,"separator":"; "},"spelling":"american"}
        self.data("08_submission/target_journal.json",target)
        style_path=self.data("08_submission/docx_style.json",DEFAULT_STYLE); style=load_style(style_path)
        doc=Document(); doc.add_heading("Synthetic medical manuscript specimen",0); doc.add_heading("Methods",1)
        doc.add_paragraph("Colour at the centre."); doc.add_paragraph("Keywords: one, two, three")
        path=self.project/"profile.docx"; doc.save(path); normalize_docx(path,style,"tables")
        strings=[paragraph.text for paragraph in Document(path).paragraphs]
        self.assertIn("Patients and methods",strings); self.assertIn("Color at the center.",strings); self.assertIn("Keywords: one; two; three",strings)

    def test_blind_check_catches_author_in_supplement_and_metadata(self):
        self.data("00_input/author_info.json",{"authors":[{"name":"Synthetic Author","email":"test@example.org"}]})
        self.data("08_submission/target_journal.json",{"blinded_main_manuscript":True})
        path=self.project/"supp.docx"; doc=Document(); doc.core_properties.author="Synthetic Author"; doc.add_paragraph("Synthetic Author"); doc.save(path)
        self.assertTrue(journalspec.blind_problems(path,self.project,"supplementary"))
        self.assertEqual(journalspec.blind_problems(path,self.project,"title_page"),[])

    def test_registered_administrative_ids_do_not_waive_results(self):
        self.data("00_input/author_info.json",{"ethics_approval":"IRB-2026-1234"})
        before="# Results\n\nIRB-2026-1234 is not a numerical result.\n\n# Declarations and Statements\n\nIRB-2026-1234 approved the study."
        after=journalspec.admin_scrub(before,self.project,"07_manuscript/full_manuscript.md")
        self.assertEqual(after.count("IRB-2026-1234"),1)

    def test_csl_patch_is_local_and_does_not_change_library_or_shared_style(self):
        from wfcore.cslpatch import apply
        base=self.text("base.csl",'<style xmlns="http://purl.org/net/xbiblio/csl"><bibliography hanging-indent="true"/></style>')
        specification=self.data("patch.json",{"source":"https://journal.example.org/authors","changes":[{"selector":".//csl:bibliography","attributes":{"hanging-indent":"false"}}]})
        output=self.project/"derived.csl"; apply(base,specification,output)
        self.assertIn('hanging-indent="false"',output.read_text()); self.assertIn('hanging-indent="true"',base.read_text())

    def test_metadata_only_rebuild_does_not_invalidate_package(self):
        path=self.project/"08_submission/bundle/a.docx"; path.parent.mkdir(parents=True)
        doc=Document(); doc.add_paragraph("Synthetic source 162"); doc.save(path)
        self.data("08_submission/bundle/manifest.json",{"items":[{"file":"08_submission/bundle/a.docx","role":"test"}]})
        write_freeze(self.project)
        doc=Document(path); doc.core_properties.last_modified_by="Another local editor"; doc.save(path)
        self.assertTrue(verify_freeze(self.project)[0],verify_freeze(self.project)[1])

    def test_author_copy_requires_full_inventory_and_blocks_science_drift(self):
        path=self.project/"08_submission/bundle/a.docx"; path.parent.mkdir(parents=True)
        doc=Document(); doc.add_paragraph("Synthetic source 162"); doc.save(path); normalize_docx(path,DEFAULT_STYLE,"tables")
        self.data("08_submission/docx_style.json",DEFAULT_STYLE)
        self.data("08_submission/bundle/manifest.json",{"items":[{"file":"08_submission/bundle/a.docx","role":"tables"}]})
        write_freeze(self.project); self.text("08_submission/independent_submission_audit.md","# Synthetic audit fixture\n\n## Verdict\nPASS\n")
        finalcopy.prepare(self.project); finalcopy.record(self.project,["Manual synthetic typeface spacing adjustment"],"Test author explicitly chooses final-copy handoff")
        self.assertEqual(finalcopy.check(self.project),[])
        copy=self.project/finalcopy.DIR/"a.docx"; doc=Document(copy); doc.paragraphs[0].text="Synthetic source 163"; doc.save(copy)
        self.assertTrue(finalcopy.check(self.project)); self.assertTrue(finalcopy.verified(self.project))

    def test_office_trailing_bytes_are_not_treated_as_metadata(self):
        path=self.project/"a.docx"; doc=Document(); doc.add_paragraph("Synthetic fixture"); doc.save(path)
        before=content_sha256(path); path.write_bytes(path.read_bytes()+b"not approved container content")
        self.assertNotEqual(before,content_sha256(path))

    def test_pdf_normalization_is_parseable_and_idempotent(self):
        from pypdf import PdfWriter, PdfReader
        path=self.project/"test.pdf"; writer=PdfWriter(); writer.add_blank_page(width=300,height=400)
        writer.add_metadata({"/CreationDate":"D:20261017101010Z"}); writer.write(path)
        normalize(path); before=path.read_bytes(); normalize(path)
        self.assertEqual(path.read_bytes(),before); self.assertEqual(len(PdfReader(path).pages),1)

    def test_sealed_presentation_batch_keeps_checkpoint_and_approval(self):
        from types import SimpleNamespace
        from rework import _cmd_batch, ROUTES
        self.freeze(); self.state.load(); self.state.data["current"]="S24_package_human_review"; self.state.save()
        plan=self.data("temp/plan.json",{"feedback_verbatim":"Correct the spelling in the already approved scientific Methods only.",
             "interpretation":"Apply a meaning-preserving spelling correction, preserving every fact and existing approval.",
             "items":[{"kind":"methods-wording","change_type":"wording","request":"Change centre to center without any scientific change.",
                       "affected_sources":["07_manuscript/methods.md"],"acceptance_criteria":["Clinical meaning, quantities and sources remain unchanged"]}]})
        args=SimpleNamespace(plan=plan,sealed=True,why=None)
        self.assertEqual(_cmd_batch(args,self.project,self.pipeline,self.state,[*ROUTES,"manuscript-copyedit","word-format-only"]),0)
        self.assertEqual(self.state.current,"S24_package_human_review")
        self.assertEqual(self.state.decision("manuscript_human_reviewed")["value"],"NO_FURTHER_REVIEW")
        record=revision.active(self.project); self.assertTrue(record["presentation_candidate"])

    def test_closed_round_accepts_only_known_manifest_metadata_drift(self):
        manifest=self.data("04_tables/manifest.json",{"created_at":"old","rows":[{"n":0,"value":162}]})
        from wfcore.determinism import receipt
        row=receipt(manifest,"04_tables/manifest.json")
        record={"schema_version":2,"status":"complete","validation":{"ok":True,"input_signature":"synthetic","results":[{"ok":True}]},
                "items":[{"id":"r1","status":"done","validated_by":["synthetic"],"changed_files":[row]}]}
        self.data(".wf/revisions/R001.json",record)
        ctx=Ctx(self.pipeline,self.state,self.project,self.pipeline.stage("S24_package_human_review"),{})
        self.data("04_tables/manifest.json",{"created_at":"new","rows":[{"n":0,"value":162}]})
        self.assertTrue(get("revision_rounds_closed")(ctx).ok)
        self.data("04_tables/manifest.json",{"created_at":"new","rows":[{"n":0,"value":163}]})
        self.assertFalse(get("revision_rounds_closed")(ctx).ok)

    def test_package_amendment_creates_new_release_and_preserves_old_release(self):
        path=self.project/"08_submission/bundle/a.docx"; path.parent.mkdir(parents=True)
        doc=Document(); doc.add_paragraph("Synthetic source 162"); doc.save(path)
        self.data("08_submission/bundle/manifest.json",{"items":[{"file":"08_submission/bundle/a.docx","role":"tables"}]})
        freeze_path=write_freeze(self.project); parent=json.loads(freeze_path.read_text()); original=self.project/parent["release_dir"]/"a.docx"; before=original.read_bytes()
        record=self.round(); doc=Document(path); doc.paragraphs[0].runs[0].bold=True; doc.save(path)
        self.assertEqual(delta.classify(self.project,record)["class"],"presentation")
        record.update(status="complete",validation={"ok":True},content_classification={"class":"presentation"})
        self.data(".wf/revisions/R001.json",record); delta.commit(self.project,record,"Synthetic layout amendment preserving exact scientific content")
        amended=json.loads(freeze_path.read_text()); self.assertNotEqual(amended["freeze_id"],parent["freeze_id"])
        self.assertEqual(original.read_bytes(),before); self.assertTrue(verify_freeze(self.project)[0])


if __name__ == "__main__": unittest.main(verbosity=2)
