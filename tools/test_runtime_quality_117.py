"""Synthetic regression cases derived from runtime retrospectives; no patient/manuscript access."""
import argparse
import base64
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from wfcore import registry,argtext
from wfcore.state import State
from wfcore.checks import Ctx
from wfcore.readiness import atomic_json

ROOT=Path(__file__).resolve().parents[1]


class RuntimeQualityTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory(prefix="medpaper-runtime-quality-"); self.addCleanup(self.temp.cleanup)
        self.project=Path(self.temp.name); self.pipe=registry.load()
        self.state=State(self.project).create("medpaper","1.17.0","S19_human_review")
    def text(self,rel,value):
        path=self.project/rel; path.parent.mkdir(parents=True,exist_ok=True); path.write_text(value,encoding="utf-8"); return path
    def data(self,rel,value): atomic_json(self.project/rel,value)
    def ctx(self,name,stage="S10_tables",**spec):
        return Ctx(self.pipe,self.state,self.project,self.pipe.stage(stage),{"check":name,**spec})
    def cli(self,*args):
        return subprocess.run([sys.executable,*args],cwd=ROOT,env={**os.environ,"MEDPAPER_PROJECT":str(self.project),"MEDPAPER_ROOT":str(ROOT),"PYTHONIOENCODING":"utf-8"},capture_output=True,text=True,encoding="utf-8")
    def test_author_text_file_and_base64_preserve_actual_words(self):
        words="我确认这一版\r\n"; path=self.project/"words.txt"; path.write_bytes(b"\xef\xbb\xbf"+words.encode("utf-8"))
        parser=argparse.ArgumentParser(); argtext.add(parser,"--words",required=True)
        self.assertEqual(parser.parse_args(["--words-file",str(path)]).words,words)
        self.assertEqual(parser.parse_args(["--words-b64",base64.b64encode(words.encode()).decode()]).words,words)
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit): parser.parse_args(["--words","text","--words-file",str(path)])
    def test_corrupt_text_is_rejected_without_guessed_repair(self):
        for value in ("我确认\ufffd一版","actual\x00words"):
            with self.assertRaises(argparse.ArgumentTypeError): argtext.intact(value)
        with self.assertRaises(argparse.ArgumentTypeError): argtext.from_b64("not base64!")
    def test_historical_owner_check_does_not_reject_current_review_files(self):
        from wfcore.checks.fsx import single_section_written
        self.state.stage_info("S08_methods")["status"]="done"
        self.text("07_manuscript/human_review.md","Synthetic current review content.")
        self.assertTrue(single_section_written(self.ctx("single_section_written","S08_methods")).ok)
        self.state.data["current"]="S08_methods"
        self.assertFalse(single_section_written(self.ctx("single_section_written","S08_methods")).ok)
    def test_office_lock_globs_exclude_only_temporary_office_files(self):
        self.text("04_tables/main/~$Table1.xlsx","Office lock")
        self.text("04_tables/main/Table1.xlsx","Actual workbook path only")
        self.assertEqual([p.name for p in self.ctx("numbers_have_provenance").glob("04_tables/main/*.xlsx")],["Table1.xlsx"])
    def test_caption_needs_actual_content_not_magic_footnote_keyword(self):
        from wfcore.checks.artifacts import legends_cover_plan
        self.data("01_protocol/artifact_plan.json",{"main_tables":[{"id":"Table 1"}]})
        self.text("05_figures/legends.md","# Figure legends\n")
        self.text("04_tables/table_captions.md","# Table captions\n\n## Table 1. Synthetic participant characteristics.\n\nValues describe the eligible study population using the stated units.\n")
        self.assertTrue(legends_cover_plan(self.ctx("legends_cover_plan")).ok)
    def make_table(self,notes=None,header=None):
        from tables.threeline import write_table
        path=self.project/"04_tables/main/Table1.xlsx"
        write_table(path=path,sheet="Table 1",title="Table 1. Synthetic participant characteristics.",header=header or ["Characteristic","Value"],rows=[["Clinical label","Not estimated"]],footnotes=notes or ["Synthetic decoding information for this controlled test."])
        return path
    def test_real_footnotes_allow_long_coherent_notes_and_reject_missing_notes(self):
        from wfcore import xlsxlite
        from wfcore.checks.artifacts import _audit_sheet
        with self.assertWarns(UserWarning): path=self.make_table(["Synthetic decoding note. "*80])
        book=xlsxlite.Workbook(path); self.assertEqual(_audit_sheet(book,book.sheets[0]),[])
        from openpyxl import load_workbook
        wb=load_workbook(path); wb.active.delete_rows(wb.active.max_row); wb.save(path); wb.close()
        book=xlsxlite.Workbook(path); self.assertTrue(any("no footnote" in p for p in _audit_sheet(book,book.sheets[0])))
    def test_unit_constant_is_bound_to_real_header_and_cannot_exempt_a_result(self):
        from wfcore.claims import table_unit_constants,validate_tables
        path=self.make_table(header=["Rate per 1000 reports","Value"])
        source="Event rates are expressed per 1000 reports in this synthetic analysis."
        self.text("01_protocol/protocol.md",source)
        row={"kind":"unit_scale","path":"04_tables/main/Table1.xlsx","sheet":"Table 1","cell":"A2","text":"Rate per 1000 reports","source_path":"01_protocol/protocol.md","source_quote":source,"reason":"This is the declared rate scaling unit, never an observed result count."}
        self.data("07_manuscript/claim_bindings.json",{"schema_version":1,"claims":[row]})
        self.assertEqual(table_unit_constants(self.project,path,"Table 1","A2",row["text"]),{"1000"})
        self.assertEqual(validate_tables(self.project,[path]),[])
        from openpyxl import load_workbook
        wb=load_workbook(path); wb.active["A3"]=row["text"]; wb.save(path); wb.close()
        row["cell"]="A3"; self.data("07_manuscript/claim_bindings.json",{"schema_version":1,"claims":[row]})
        with self.assertRaises(ValueError): table_unit_constants(self.project,path,"Table 1","A3",row["text"])
    def test_workbook_permission_failure_is_pending_io_and_never_a_provenance_pass(self):
        from wfcore.checks.numbers import numbers_have_provenance
        self.make_table(); self.data("03_analysis/results/synthetic.json",{"n":1})
        with patch("wfcore.xlsxlite.numeric_cell_values",side_effect=PermissionError("Synthetic Excel lock")):
            result=numbers_have_provenance(self.ctx("numbers_have_provenance",source="tables"))
        self.assertFalse(result.ok); self.assertIn("IO pending",result.detail); self.assertIn("Excel",result.hints[0])
    def test_unit_binding_proof_change_invalidates_the_cached_gate(self):
        from wfcore.gatecache import Receipts
        source=self.text("03_analysis/code/rate.R","# Synthetic source says rates per 1000 reports\n")
        self.data("07_manuscript/claim_bindings.json",{"schema_version":1,"claims":[{"kind":"unit_scale","source_path":"03_analysis/code/rate.R"}]})
        receipts=Receipts(self.project,self.pipe,self.state)
        stage=self.pipe.stage("S10_tables"); spec={"check":"numbers_have_provenance","source":"tables"}
        first=receipts.identity(stage,spec)
        source.write_text("# Source now uses per 100 reports\n",encoding="utf-8")
        self.assertNotEqual(first,receipts.identity(stage,spec))
    def test_stacks_in_both_orientations_pass_but_floating_or_truncated_bars_fail(self):
        import matplotlib; matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        from figures.elements import baseline_zero
        for horizontal in (True,False):
            fig,ax=plt.subplots()
            if horizontal: ax.barh([0,1],[3,4]); ax.barh([0,1],[2,3],left=[3,4])
            else: ax.bar([0,1],[3,4]); ax.bar([0,1],[2,3],bottom=[3,4])
            self.assertTrue(baseline_zero(fig)[0]); plt.close(fig)
        for truncated in (True,False):
            fig,ax=plt.subplots(); ax.bar([0,1],[3,4],bottom=0 if truncated else 1)
            if truncated: ax.set_ylim(2,8)
            self.assertFalse(baseline_zero(fig)[0]); plt.close(fig)
    def test_renderer_and_qc_cannot_run_on_same_figure_and_old_report_survives(self):
        from wfcore.figurejobs import job,marker,FigureBusy
        from figures.qc import qc_figure
        path=self.text("05_figures/qc/qc_report.json","{\"old\":true}"); before=path.read_bytes()
        self.data("01_protocol/artifact_plan.json",{"main_figures":[{"id":"Figure 1","file":"05_figures/out/Figure1.png"}]})
        with job(self.project,"Figure 1","render"):
            result=qc_figure({"id":"Figure 1","file":"05_figures/out/Figure1.png"},{},project=self.project)
            self.assertTrue(result["pending"]); self.assertFalse(result["ok"])
            command=self.cli("tools/figures/qc.py","--all")
            self.assertEqual(command.returncode,3,command.stdout+command.stderr)
            self.assertIn("PENDING",command.stdout)
            from wfcore.checks.artifacts import figures_qc_pass
            gate=figures_qc_pass(self.ctx("figures_qc_pass","S11_figures"))
            self.assertFalse(gate.ok); self.assertIn("pending",gate.detail)
            with self.assertRaises(FigureBusy):
                with job(self.project,"Figure 1","QC"): pass
            with job(self.project,"Figure 2","render"): pass
        self.assertEqual(path.read_bytes(),before); self.assertFalse(marker(self.project,"Figure 1").exists())
    def test_actual_r_stacked_positive_negative_and_horizontal_bars(self):
        import shutil
        rscript=shutil.which("Rscript")
        if not rscript: self.skipTest("Rscript unavailable")
        script=self.text("stacked.R",'''grDevices::pdf(file=NULL)
source("tools/figures/r_style.R")
library(ggplot2)
d <- data.frame(g=rep(c("First", "Second"), each=3), part=rep(c("Mild", "Moderate", "Severe"), 2), n=c(3, 2, -1, 4, 2, -2))
v <- ggplot(d, aes(g, n, fill=part)) + geom_col()
h <- ggplot(d, aes(n, g, fill=part)) + geom_col()
stopifnot(medpaper_layers(v)$bar_baseline_zero, medpaper_layers(h)$bar_baseline_zero)
stopifnot(!medpaper_layers(v + coord_cartesian(ylim=c(2,8)))$bar_baseline_zero)
''')
        result=subprocess.run([rscript,"--vanilla","--encoding=UTF-8",str(script)],cwd=ROOT,capture_output=True,text=True,encoding="utf-8",timeout=30)
        self.assertEqual(result.returncode,0,result.stdout+result.stderr)
    def test_failed_figure_job_removes_only_its_own_marker(self):
        from wfcore.figurejobs import job,marker
        with self.assertRaises(RuntimeError):
            with job(self.project,"Figure 1","render"): raise RuntimeError("Synthetic failure")
        self.assertFalse(marker(self.project,"Figure 1").exists())
    def test_heartbeat_keeps_stdout_machine_readable_and_real_failures_propagate(self):
        from wfcore.execution import run_logged
        stderr=io.StringIO()
        with contextlib.redirect_stderr(stderr): result=run_logged([sys.executable,"-c","import time; time.sleep(.15); print('{\"actual\":true}')"],project=self.project,heartbeat=.05,timeout=2)
        self.assertEqual(json.loads(result.stdout),{"actual":True}); self.assertIn("Completion is pending",stderr.getvalue())
        result=run_logged([sys.executable,"-c","raise SystemExit(7)"],project=self.project,timeout=2)
        self.assertEqual(result.returncode,7)
    def test_reference_count_show_before_title_page_is_read_only(self):
        self.text("07_manuscript/full_manuscript.md","# Synthetic paper\n\nReal citation syntax only [@synthetic1; @synthetic2].\n")
        result=self.cli("tools/manuscript/reference_count.py","show","--project",str(self.project))
        self.assertEqual(result.returncode,0,result.stderr); self.assertIn("2",result.stdout)
        self.assertFalse((self.project/"08_submission/integration/title_page.md").exists())
    def test_handoff_is_bounded_and_version_refresh_does_not_rewind_modern_tail(self):
        self.state.data["pipeline_version"]="1.16.0"; self.state.data["current"]="S24_package_human_review"; self.state.save()
        result=self.cli("tools/wf.py","handoff")
        self.assertEqual(result.returncode,0,result.stderr); self.assertLessEqual(len(result.stdout.splitlines()),60)
        current=State(self.project).load(); self.assertEqual(current.current,"S24_package_human_review")
        self.assertEqual(current.data["engine_pin"],"1.17.0")
    def test_compact_cards_retain_the_complete_procedure_via_real_cli(self):
        for name in ("S07_artifacts","S20_journal","S23_package","S25_submission_audit"):
            card=ROOT/"pipeline/stages"/(name+".md")
            detail=ROOT/"reference/stage_details"/card.name
            self.assertLess(len(card.read_text(encoding="utf-8")),4000)
            self.assertTrue(detail.is_file())
            result=self.cli("tools/wf.py","card",name,"--details")
            self.assertEqual(result.returncode,0,result.stderr)
            self.assertEqual(result.stdout.strip(),detail.read_text(encoding="utf-8").strip())
    def test_line_numbers_are_journal_driven_and_audited_in_real_xml(self):
        from docx import Document
        from wfcore.journalspec import apply_docx,docx_problems
        doc=Document(); doc.add_paragraph("Synthetic document for line-number XML testing.")
        target={"line_numbers":True}; apply_docx(doc,target,"manuscript")
        path=self.project/"synthetic.docx"; doc.save(path)
        self.assertFalse(docx_problems(path,target,self.project,"manuscript"))
        target["line_numbers"]=False
        self.assertTrue(docx_problems(path,target,self.project,"manuscript"))
        apply_docx(doc,target,"manuscript"); doc.save(path)
        self.assertFalse(docx_problems(path,target,self.project,"manuscript"))


if __name__=="__main__": unittest.main(verbosity=2)
