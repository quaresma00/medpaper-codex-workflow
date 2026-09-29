"""Synthetic behavioral regressions: quality gates and bounded reuse, no real studies."""
import argparse
import contextlib
import copy
import io
import json
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

from wfcore import claims, dependencies, gates, registry, revision
from wfcore.checks import Ctx, Result, get, load_all
from wfcore.checks.numbers import _offenders, numbers_have_provenance
from wfcore.checks.revisions import revision_rounds_closed
from wfcore.packagefreeze import _sha256
from wfcore.readiness import atomic_json
from wfcore.state import State
from rework import _cmd_close, _cmd_mark


class IncrementalQualityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="medpaper-incremental-")
        self.addCleanup(self.tmp.cleanup)
        self.project = Path(self.tmp.name)
        self.pipe = registry.load()
        self.state = State(self.project).create("medpaper", self.pipe.meta["version"], "S19_human_review")
        load_all()

    def text(self, rel, value):
        path = self.project / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(value, encoding="utf-8")
        return path

    def data(self, rel, value):
        atomic_json(self.project / rel, value)

    def ctx(self, check, **spec):
        return Ctx(self.pipe, self.state, self.project, self.pipe.stage("S19_human_review"), {"check": check, **spec})

    def facts(self):
        context = {"outcome": "mortality", "population": "eligible cohort", "comparison": "exposed versus control",
                   "model": "adjusted model", "timepoint": "one year", "unit": "risk ratio", "denominator": "eligible patients"}
        self.data("03_analysis/results/primary.json", {"reported": {"context": context,
                  "values": {"estimate": 1.34, "p_value": 0.23, "n": 123}}})
        self.data("01_protocol/study_facts.json", {"numeric_facts": {"mortality": {
            "source": "03_analysis/results/primary.json", "pointer": "/reported", "context": context}}})

    def prose(self, text, ids=None):
        self.text("07_manuscript/results.md", "# Results\n\n" + text)
        self.data("07_manuscript/claim_bindings.json", {"schema_version": 1, "claims": [{
            "path": "07_manuscript/results.md", "text": text, "fact_ids": ids or ["mortality"]}]})
        return numbers_have_provenance(self.ctx("numbers_have_provenance", path="07_manuscript/results.md"))

    def test_valid_specific_finding(self):
        self.facts()
        self.assertTrue(self.prose("The mortality risk ratio was 1.34 (P=0.23).").ok)

    def test_wrong_sign_rejected(self):
        self.facts()
        self.assertFalse(self.prose("The mortality risk ratio was -1.34.").ok)
        self.assertTrue(_offenders("Estimate −1.34", {1.34}, {"1.34"}, set()))

    def test_wrong_endpoint_rejected_even_with_same_number(self):
        self.facts()
        self.assertFalse(self.prose("The stroke risk ratio was 1.34.").ok)

    def test_unsupported_p_threshold_rejected(self):
        self.facts()
        self.assertFalse(self.prose("The mortality risk ratio was 1.34 (P<0.001).").ok)

    def test_small_invented_number_no_longer_exempt(self):
        self.facts()
        self.assertFalse(self.prose("The mortality risk ratio was 4.").ok)

    def test_context_drift_rejected(self):
        self.facts()
        source = self.project / "03_analysis/results/primary.json"
        payload = json.loads(source.read_text())
        payload["reported"]["context"]["model"] = "different model"
        self.data("03_analysis/results/primary.json", payload)
        self.assertFalse(self.prose("The mortality risk ratio was 1.34.").ok)

    def test_exact_prose_binding_stale_or_missing_fails(self):
        self.facts()
        self.prose("The mortality risk ratio was 1.34.")
        self.text("07_manuscript/results.md", "The mortality risk ratio was 1.34 in 123 participants.")
        result = numbers_have_provenance(self.ctx("numbers_have_provenance", path="07_manuscript/results.md"))
        self.assertFalse(result.ok)

    def test_ci_and_scientific_notation(self):
        self.assertFalse(_offenders("HR 1.34 (95% CI 1.20 to 1.50)", {1.34, 1.2, 1.5}, set(), set()))
        self.assertFalse(_offenders("Estimate -1.2e-3", {-0.00123}, set(), set()))
        self.assertTrue(_offenders("Estimate -1.2e-3", {0.00123}, set(), set()))

    def test_table_cell_binding_and_value_drift(self):
        from openpyxl import Workbook
        self.facts()
        path = self.project / "04_tables/main/Table1.xlsx"
        path.parent.mkdir(parents=True)
        wb = Workbook(); ws = wb.active; ws.title = "Table 1"
        ws.append(["Outcome", "Risk ratio"]); ws.append(["mortality", 1.34]); wb.save(path)
        self.data("07_manuscript/claim_bindings.json", {"schema_version": 1, "claims": [{
            "path": "04_tables/main/Table1.xlsx", "sheet": "Table 1", "cell": "B2",
            "text": "1.34", "label_cells": ["A2"], "fact_ids": ["mortality"]}]})
        self.assertFalse(claims.validate_tables(self.project, [path]))
        ws["A2"] = "stroke"; wb.save(path)
        self.assertTrue(claims.validate_tables(self.project, [path]))
        ws["A2"] = "mortality"; ws["B2"] = -1.34; wb.save(path)
        self.assertTrue(claims.validate_tables(self.project, [path]))

    def test_assembly_and_journal_bindings_remain_backstage(self):
        self.facts()
        sentence = "The mortality risk ratio was 1.34."
        self.prose(sentence)
        manuscript = "# Abstract\n\nNo numerical claims.\n\n# Results\n\n" + sentence
        self.data("07_manuscript/claim_bindings.json", claims.assembled_bindings(self.project, manuscript))
        self.assertFalse(claims.validate_prose(self.project, "07_manuscript/full_manuscript.md", manuscript))
        claims.copy_for_journal(self.project)
        self.assertFalse(claims.validate_prose(self.project, "08_submission/integration/full_manuscript.md", manuscript))
        path = self.project / "08_submission/integration/claim_bindings.json"
        before = path.read_bytes(); claims.copy_for_journal(self.project)
        self.assertEqual(before, path.read_bytes())

    def simple_stage(self):
        stage = copy.deepcopy(self.pipe.stage("S09_results"))
        stage.gate = [{"check": "md_sections", "path": "07_manuscript/results.md", "headings": ["Results"]}]
        self.text("07_manuscript/results.md", "# Results\n\nA qualitative finding.")
        return stage

    def test_unchanged_pure_gate_reused_but_changed_source_rechecked(self):
        stage = self.simple_stage()
        original = get("md_sections")
        with patch("wfcore.checks.text.md_sections", wraps=original):
            gates.run_stage(self.pipe, self.state, self.project, stage, record=True)
        result = gates.run_stage(self.pipe, self.state, self.project, stage, reuse=True)
        self.assertTrue(result[0].ok); self.assertEqual(self.state.data["check_reuse"]["reused"], 1)
        self.text("07_manuscript/results.md", "# Wrong heading")
        result = gates.run_stage(self.pipe, self.state, self.project, stage, reuse=True)
        self.assertFalse(result[0].ok); self.assertEqual(self.state.data["check_reuse"]["reused"], 0)

    def test_expired_receipt_and_config_change_recheck(self):
        stage = self.simple_stage()
        gates.run_stage(self.pipe, self.state, self.project, stage, record=True)
        path = self.state.dir / "check_receipts.json"
        rows = json.loads(path.read_text())
        for row in rows.values():
            row["at"] = time.time() - 121
        atomic_json(path, rows)
        gates.run_stage(self.pipe, self.state, self.project, stage, reuse=True)
        self.assertEqual(self.state.data["check_reuse"]["reused"], 0)
        self.state.record_decision("review", "YES", "New explicit decision")
        result = gates.run_stage(self.pipe, self.state, self.project, stage, reuse=True)
        self.assertTrue(result[0].ok)
        self.assertEqual(self.state.data["check_reuse"]["reused"], 0)

    def test_default_manuscript_path_is_bound_in_cache(self):
        from wfcore.gatecache import Receipts
        stage = self.pipe.stage("S17_assemble")
        spec = {"check": "manuscript_structure"}
        self.text("07_manuscript/full_manuscript.md", "Initial manuscript")
        cache = Receipts(self.project, self.pipe, self.state)
        before = cache.identity(stage, spec)
        self.text("07_manuscript/full_manuscript.md", "Changed manuscript")
        self.assertNotEqual(before, cache.identity(stage, spec))

    def test_live_and_unknown_checks_never_cached(self):
        stage = self.simple_stage()
        stage.gate = [{"check": "reference_provenance", "live": True}]
        with patch("wfcore.checks.get", return_value=lambda ctx: Result(True, "reference_provenance")) as called:
            gates.run_stage(self.pipe, self.state, self.project, stage, record=True)
            gates.run_stage(self.pipe, self.state, self.project, stage, reuse=True)
            self.assertEqual(called.call_count, 2)
        self.assertEqual(self.state.data["check_reuse"]["reused"], 0)

    def test_missing_inputs_cannot_force_past_unexecuted_integrity_gates(self):
        from wfcore import cli
        stage = copy.deepcopy(self.pipe.stage("S13_reflib"))
        results = gates.run_stage(self.pipe, self.state, self.project, stage, reuse=True)
        self.assertTrue(any(r.check == "dependent_checks_pending" and r.blocking for r in results))
        self.state.data["current"] = stage.id
        with patch.object(cli, "_load", return_value=(self.pipe, self.state, self.project)), contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cli.cmd_advance(argparse.Namespace(force=True, note="Attempt an invalid bypass")), 2)
        self.assertEqual(self.state.current, stage.id)

    def test_entry_only_plan_diff_preserves_other_display(self):
        plan = {"main_figures": [{"id": "Figure 1", "file": "05_figures/out/Figure1.png", "width": "single"},
                                 {"id": "Figure 2", "file": "05_figures/out/Figure2.png", "width": "single"}]}
        old = copy.deepcopy(plan); plan["main_figures"][1]["width"] = "double"
        self.data("01_protocol/artifact_plan.json", plan)
        result = dependencies.closure(self.project, ["01_protocol/artifact_plan.json"],
                                      baseline_documents={"01_protocol/artifact_plan.json": old})
        self.assertIn("05_figures/out/Figure2.png", result)
        self.assertNotIn("05_figures/out/Figure1.png", result)
        plan["shared_style"] = "new shared setting"; self.data("01_protocol/artifact_plan.json", plan)
        result = dependencies.closure(self.project, ["01_protocol/artifact_plan.json"],
                                      baseline_documents={"01_protocol/artifact_plan.json": old})
        self.assertIn("05_figures/out/Figure1.png", result)

    def test_missing_dependency_snapshot_fails_conservative(self):
        self.data("01_protocol/artifact_plan.json", {"main_figures": [{"id": "Figure 1", "file": "05_figures/out/Figure1.png"}]})
        result = dependencies.closure(self.project, ["01_protocol/artifact_plan.json"])
        self.assertIn("05_figures/out/Figure1.png", result)

    def test_two_closed_rounds_with_evidence_refresh_and_later_text_drift(self):
        rel = "08_submission/integration/title_page.md"
        self.state.data["current"] = "S24_package_human_review"; self.state.save()
        self.text(rel, "# Title page\n\nInitial address.")
        for number in (1, 2):
            rid = f"R{number:03d}"
            row = {"schema_version": 2, "round_id": rid, "status": "active", "review_stage": self.state.current,
                   "allowed_files": [rel, "06_refs/verified.json"], "baseline": revision.snapshot(self.project),
                   "items": [{"id": rid + "-01", "kind": "title-page-or-statements", "affected_sources": [rel], "status": "pending"}],
                   "checks": [{"stage": "S21_authors", "spec": {"check": "md_sections", "path": rel, "headings": ["Title page"]}}]}
            self.state.data["active_revision_round"] = rid; self.state.save()
            self.text(rel, f"# Title page\n\nCorrected address round {number}.")
            path = self.state.dir / "revisions" / (rid + ".json")
            with contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(revision.check_round(self.project, self.pipe, self.state, path, row), 0)
                signature = revision.input_signature(self.project, row)
                # This test exercises closure bookkeeping, not PubMed truth verification.
                self.data("06_refs/verified.json", {"synthetic_refresh_counter": number})
                self.assertEqual(signature, revision.input_signature(self.project, row))
                self.assertFalse(revision.scope_problems(self.project, row))
                _cmd_mark(argparse.Namespace(item=rid + "-01", changed_file=[rel], validated_by=["actual md_sections"],
                    summary="Corrected the supplied administrative address and verified its heading."), self.project, self.state)
                _cmd_close(argparse.Namespace(summary="Completed this targeted correction and verified the final document source without changing other content."), self.state)
            self.assertTrue(revision_rounds_closed(self.ctx("revision_rounds_closed")).ok)
        self.text(rel, "# Title page\n\nUnrecorded third edit.")
        self.assertFalse(revision_rounds_closed(self.ctx("revision_rounds_closed")).ok)

    def test_invalid_renewed_evidence_cannot_close(self):
        row = {"validation": {"renewable_evidence": {}}, "checks": [{"stage": "S13_reflib", "spec": {"check": "reference_provenance", "live": True}}]}
        self.data("06_refs/verified.json", {"synthetic_invalid": True})
        with self.assertRaises(ValueError):
            revision.refresh_evidence(self.project, self.pipe, self.state, row)


if __name__ == "__main__":
    unittest.main(verbosity=2)
