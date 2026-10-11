"""Synthetic source/dependency/version regressions, never a real study or submission."""
import argparse
import contextlib
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from wfcore import dependencies,registry,revision
from wfcore.readiness import atomic_json
from wfcore.packagefreeze import _sha256
from wfcore.packagecontent import write_baseline,verify_baseline,build_baseline
from wfcore.buildproof import content_digest
from wfcore.checks import Ctx,get,load_all
from wfcore.state import State
from wfcore.journalrules import limits,compare
from wfcore.presentation import check as check_migration,MAP
from wfcore.execution import run_logged
from wfcore.submission import docx_sections
from manuscript.build_docx import normalize_docx,planned_figure_ids,audit_docx
from text.polish import facts,cmd_snapshot
from rework import _cmd_mark

DEFAULT_STYLE={"font_family":"Times New Roman","body_font_pt":12,"title_font_pt":14,
               "section_heading_font_pt":12,"subsection_heading_font_pt":12,"line_spacing":2,
               "paragraph_spacing_after_pt":0,"margins_in":1,"paper_size":"A4",
               "all_text_black":True,"external_hyperlinks":False}


class SourceAlignedTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory(prefix="medpaper-source-aligned-");self.addCleanup(self.tmp.cleanup)
        self.project=Path(self.tmp.name);self.pipe=registry.load();load_all()
        self.state=State(self.project).create("medpaper",self.pipe.meta["version"],"S24_package_human_review")

    def text(self,rel,text):
        p=self.project/rel;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(text,encoding="utf-8");return p

    def data(self,rel,value):
        atomic_json(self.project/rel,value);return self.project/rel

    def docx(self,rel="08_submission/bundle/combined.docx"):
        p=self.project/rel;p.parent.mkdir(parents=True,exist_ok=True)
        doc=Document();doc.add_paragraph("Source aligned validation specimen");doc.add_paragraph("A complete, test-only document.");doc.save(p)
        return p

    def round(self,sources,kind="cover-letter-or-package-structure",change="administrative"):
        self.state.data["active_revision_round"]="R001";self.state.save()
        row={"schema_version":2,"round_id":"R001","status":"active","review_stage":"S24_package_human_review",
             "items":[{"id":"R001-01","kind":kind,"owning_stage":"S23_package","affected_sources":sources,
                       "change_type":change,"status":"pending","acceptance_criteria":["Actual final content matches request"]}],
             "baseline":revision.snapshot(self.project),"baseline_documents":{},"feedback_updates":[{"feedback_verbatim":"Original user feedback preserved"}]}
        revision.plan_scope(self.project,self.pipe,row);path=self.data(".wf/revisions/R001.json",row)
        return path,row

    def baseline(self):
        source=self.text("08_submission/integration/combined_upload.md","Source aligned validation specimen\nA complete, test-only document.")
        second=self.text("08_submission/integration/additional.md","Additional source input")
        output=self.docx()
        revision.record_build(self.project,output,[source,second],entrypoints=[source,second],tool="actual test producer",parameters={"kind":"test"})
        self.data("08_submission/bundle/manifest.json",{"items":[{"role":"supplementary","file":output.relative_to(self.project).as_posix(),"source_files":[source.relative_to(self.project).as_posix()]}]})
        write_baseline(self.project)
        return source,second,output

    def test_combined_entrypoint_not_guessed_from_role(self):
        source,_,_=self.baseline();record=build_baseline(self.project)["records"][0]
        self.assertEqual(record["source"],source.relative_to(self.project).as_posix())
        self.assertEqual(len(record["build"]["entrypoints"]),2)
        self.assertTrue(verify_baseline(self.project)[0])

    def test_actual_secondary_input_drift_is_detected(self):
        _,second,_=self.baseline();second.write_text("Modified upstream input")
        self.assertFalse(verify_baseline(self.project)[0])

    def test_layout_whitespace_allowed_but_header_text_drift_rejected(self):
        _,_,output=self.baseline();doc=Document(output);doc.add_paragraph("");doc.save(output)
        self.assertTrue(verify_baseline(self.project)[0])
        doc=Document(output);doc.sections[0].header.paragraphs[0].text="Different patient cohort";doc.save(output)
        self.assertFalse(verify_baseline(self.project)[0])

    def test_observed_dependency_contains_composed_input_and_portal(self):
        source,_,output=self.baseline();manifest=json.loads((self.project/"08_submission/bundle/manifest.json").read_text())
        manifest["items"][0]["role"]="manuscript";self.data("08_submission/bundle/manifest.json",manifest)
        found=dependencies.closure(self.project,[source.relative_to(self.project).as_posix()],change_type="wording")
        self.assertIn(output.relative_to(self.project).as_posix(),found);self.assertIn("08_submission/portal_fields.json",found)

    def test_discovered_edge_extends_scope_without_rewriting_feedback(self):
        source="08_submission/integration/combined.md";self.text(source,"Original source")
        path,row=self.round([source]);feedback=row["feedback_updates"][:]
        revision.add_dependency(self.project,self.pipe,path,row,source,"08_submission/integration/combined_preview.md",
                                "This preview directly consumes the already affected combined manuscript source.")
        self.assertIn("08_submission/integration/combined_preview.md",row["allowed_files"])
        self.assertEqual(feedback,row["feedback_updates"])
        with self.assertRaises(ValueError):revision.add_dependency(self.project,self.pipe,path,row,source,"07_manuscript/methods.md","No journal task may authorize a frozen scientific source edit.")

    def test_nonscientific_dependency_cannot_authorize_analysis(self):
        source="07_manuscript/methods.md";self.text(source,"Methods")
        path,row=self.round([source]);row["review_stage"]="S19_human_review";revision.plan_scope(self.project,self.pipe,row)
        with self.assertRaises(ValueError):revision.add_dependency(self.project,self.pipe,path,row,source,"03_analysis/results/new.json","Presentation work must not silently authorize changed scientific estimates.")

    def test_checked_and_changed_sources_are_distinct(self):
        a="08_submission/integration/a.md";b="08_submission/integration/b.md";self.text(a,"Original a");self.text(b,"Unchanged b")
        _,row=self.round([a,b]);self.text(a,"Updated a")
        args=argparse.Namespace(item="R001-01",changed_file=[a],checked_file=[b],validated_by=["Actual verification"],summary="Updated only a and inspected unchanged b against the request.")
        self.assertEqual(_cmd_mark(args,self.project,self.state),0);item=revision.active(self.project)["items"][0]
        self.assertEqual([r["path"] for r in item["changed_files"]],[a]);self.assertEqual([r["path"] for r in item["checked_files"]],[b])

    def test_word_layout_does_not_schedule_data_acquisition(self):
        source="08_submission/bundle/title.docx";self.docx(source)
        _,row=self.round([source],kind="word-format-only",change="layout")
        self.assertNotIn("data_acquisition_complete",[r["spec"]["check"] for r in row["checks"]])
        self.assertNotIn("S05_analysis",row["validation_stages"])

    def test_exact_round_check_reused_and_drift_invalidates(self):
        source="08_submission/integration/note.md";self.text(source,"# Note\nCorrect text")
        path,row=self.round([source]);row["checks"]=[{"stage":"S23_package","spec":{"check":"md_sections","path":source,"headings":["Note"]}}];self.data(".wf/revisions/R001.json",row)
        self.assertEqual(revision.check_round(self.project,self.pipe,self.state,path,row),0)
        with patch("wfcore.revision.validate",side_effect=AssertionError("Repeated unchanged check")):
            self.assertEqual(revision.check_round(self.project,self.pipe,self.state,path,row),0)
        self.text(source,"# Wrong heading\nDrifted text")
        self.assertNotEqual(revision.check_round(self.project,self.pipe,self.state,path,row),0)

    def migration(self):
        rel="08_submission/integration/full_manuscript.md";supp="08_submission/integration/supplementary_methods.md"
        old="Approved method used 20 observations with P = 0.04 [@fixturekey]. Figure 1 shows the model."
        self.text(rel,old);self.text(supp,"Supplementary methods.")
        snap=self.project/"08_submission/integration/prepolish";snap.mkdir()
        (snap/"full_manuscript.md").write_text(old);(snap/"supplementary_methods.md").write_text("Supplementary methods.")
        index={"snapshot_id":"test-snapshot","files":{rel:facts(old),supp:facts("Supplementary methods.")}}
        mapped=old.replace("Figure 1","Figure S1");self.text(rel,"Main narrative.");self.text(supp,"Supplementary methods.\n"+mapped)
        self.data(MAP,{"snapshot_id":"test-snapshot","renumbering":{"Figure 1":"Figure S1"},"moves":[{"from":rel,"to":supp,"text":old}]})
        return snap,index,rel,supp

    def test_package_movement_and_renumbering_preserve_facts(self):
        snap,index,_,_=self.migration();self.assertEqual(check_migration(self.project,snap,index,facts),[])

    def test_migration_rejects_changed_p_value_or_snapshot(self):
        snap,index,_,supp=self.migration();self.text(supp,(self.project/supp).read_text().replace("0.04","0.09"))
        with self.assertRaises(ValueError):check_migration(self.project,snap,index,facts)
        index["snapshot_id"]="unrelated-snapshot"
        with self.assertRaises(ValueError):check_migration(self.project,snap,index,facts)

    def test_unapproved_method_addition_not_a_presentation_exception(self):
        snap,index,_,supp=self.migration();mapping=json.loads((self.project/MAP).read_text());mapping["additions"]=[{"source":"03_analysis/results/invented.json","to":supp,"text":"Result 99"}]
        self.text("03_analysis/results/invented.json","Result 99");self.data(MAP,mapping)
        with self.assertRaises(ValueError):check_migration(self.project,snap,index,facts)

    def test_typed_null_and_recommendation_do_not_become_caps(self):
        self.data("08_submission/submission_requirements.json",{"article_type":"Original research","word_count":{"source":"https://journal.example/original-research","limit":None,"abstract_limit":None},
                  "limits":{"abstract":{"status":"recommended","value":2,"source":"https://journal.example/original-research"}}})
        _,rules=limits(self.project);self.assertEqual(compare({"main":10000,"abstract":100},rules),[])
        rules["abstract"]["status"]="required";self.assertTrue(compare({"abstract":100},rules))

    def test_wrong_article_type_and_unspecified_nonnull_rejected(self):
        base={"article_type":"Original research","word_count":{"source":"https://journal.example"},"limits":{"main":{"status":"required","value":2000,"source":"https://journal.example","article_type":"Review"}}}
        self.data("08_submission/submission_requirements.json",base)
        with self.assertRaises(ValueError):limits(self.project)
        base["limits"]["main"]={"status":"unspecified","value":2000,"source":"https://journal.example"};self.data("08_submission/submission_requirements.json",base)
        with self.assertRaises(ValueError):limits(self.project)

    def test_structured_abstract_subheading_not_real_methods_boundary(self):
        path=self.project/"structured.docx";doc=Document();doc.add_heading("Clinical validation specimen",0);doc.add_heading("Abstract",1);doc.add_heading("Methods",2);doc.add_paragraph("We evaluated existing records.")
        for section in ("Introduction","Methods","Results","Discussion"):doc.add_heading(section,1);doc.add_paragraph("Complete body section.")
        doc.save(path);normalize_docx(path,DEFAULT_STYLE,"title_page")
        title,sections,_=docx_sections(path);self.assertIn("We evaluated existing records.",sections["Abstract"]);self.assertIn("Complete body section",sections["Methods"])
        first=content_digest(path);normalize_docx(path,DEFAULT_STYLE,"title_page");self.assertEqual(first,content_digest(path))
        from docx.oxml.ns import qn
        checked=Document(path)
        names=[node.get(qn("w:name")) for node in checked.element.iter(qn("w:bookmarkStart"))]
        self.assertTrue(all(len(name)<=40 for name in names));self.assertEqual(len(names),len(set(names)))
        self.assertIsNotNone(checked.settings.element.find(qn("w:doNotAutoCompressPictures")))

    def test_embedded_scientific_media_is_not_pure_layout(self):
        import zipfile
        path=self.docx();first=content_digest(path)
        with zipfile.ZipFile(path,"a") as archive:
            archive.writestr("word/media/validation-image.png",b"Synthetic test media, not patient data")
        self.assertNotEqual(first,content_digest(path))

    def test_journal_plan_is_separate_from_scientific_plan(self):
        self.data("01_protocol/artifact_plan.json",{"main_figures":[{"id":"Figure 1"}]})
        self.data("08_submission/integration/artifact_plan.json",{"supp_figures":[{"id":"Figure S1"}]})
        scientific=Ctx(self.pipe,self.state,self.project,self.pipe.stage("S19_human_review"),{})
        journal=Ctx(self.pipe,self.state,self.project,self.pipe.stage("S24_package_human_review"),{})
        self.assertEqual(scientific.read_json("01_protocol/artifact_plan.json")["main_figures"][0]["id"],"Figure 1")
        self.assertEqual(journal.read_json("01_protocol/artifact_plan.json")["supp_figures"][0]["id"],"Figure S1")
        self.assertEqual(planned_figure_ids(self.project,scope="supp"),["Figure S1"])

    def test_task_process_timeout_and_failure_have_owned_logs(self):
        failed=run_logged([sys.executable,"-c","import sys;print('full failure evidence');sys.exit(3)"],project=self.project)
        self.assertEqual(failed.returncode,3);self.assertIn("full failure evidence",failed.log_path.read_text())
        with self.assertRaises(ValueError):run_logged([sys.executable,"-c","import time;time.sleep(5)"],project=self.project,timeout=.1)
        records=[json.loads(p.read_text()) for p in (self.project/".wf/processes").glob("*.json")]
        self.assertIn("timeout",{r["status"] for r in records});self.assertTrue(all(r["pid"]!=os.getpid() for r in records))

    def test_new_round_snapshot_versioned_and_same_round_cannot_bless_edits(self):
        rel="08_submission/integration/full_manuscript.md";self.text(rel,"# Methods\nApproved method used 20 records.")
        _,row=self.round([rel]);args=argparse.Namespace(force=False,new_round=False)
        with patch("text.polish.project_root",return_value=self.project),contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cmd_snapshot(args),0)
            old=json.loads((self.project/"08_submission/integration/prepolish/facts.json").read_text())
            self.text(rel,"# Methods\nChanged method used 99 records.")
            with self.assertRaises(ValueError):cmd_snapshot(argparse.Namespace(force=True,new_round=False))
            self.text(rel,"# Methods\nApproved method used 20 records.")
            row["round_id"]="R002";self.data(".wf/revisions/R002.json",row);self.state.data["active_revision_round"]="R002";self.state.save()
            self.assertEqual(cmd_snapshot(argparse.Namespace(force=False,new_round=True)),0)
            current=json.loads((self.project/"08_submission/integration/prepolish/facts.json").read_text())
            self.assertNotEqual(old["snapshot_id"],current["snapshot_id"])
            self.assertTrue((self.project/".wf/polish_snapshots"/(old["snapshot_id"]+".json")).is_file())

    def test_fact_query_preview_bounded_and_complete_catalog_preserved(self):
        context={k:"Synthetic test context" for k in ("outcome","population","comparison","model","timepoint","unit","denominator")}
        self.data("03_analysis/results/fixture.json",{"reported":{"fixture":{"context":context,"values":{"full_array":list(range(1000))}}}})
        self.data("01_protocol/study_facts.json",{"numeric_facts":{f"fact{i}":{"source":"03_analysis/results/fixture.json","pointer":"/reported/fixture","context":context} for i in range(30)}})
        tool=Path(__file__).parent/"manuscript/readiness.py"
        run=subprocess.run([sys.executable,str(tool),"facts","--project",str(self.project),"--limit","2"],capture_output=True,text=True)
        self.assertEqual(run.returncode,0,run.stderr);data=json.loads(run.stdout);self.assertEqual(data["shown"],2);self.assertLess(len(run.stdout),5000)
        full=json.loads(Path(data["complete_file"]).read_text());self.assertEqual(len(full),30);self.assertEqual(len(full["fact0"]["values"]["full_array"]),1000)

    def test_guide_prose_numbers_do_not_override_typed_selected_article_limits(self):
        self.text("08_submission/guidelines_extract.md","Review abstract maximum 10 words; original articles 5000 words; references 50.")
        self.data("08_submission/submission_requirements.json",{"article_type":"Original research","word_count":{"source":"https://journal.example/original","include_sections":["Introduction","Methods","Results","Discussion"],"limit":None,"abstract_limit":None}})
        self.text("08_submission/integration/full_manuscript.md","\n\n".join("# "+name+"\n"+"actual words "*20 for name in ("Abstract","Introduction","Methods","Results","Discussion")))
        result=get("journal_limits_met")(Ctx(self.pipe,self.state,self.project,self.pipe.stage("S22_polish"),{}))
        self.assertTrue(result.ok,result.detail)

    def test_actual_word_numbering_delimiter_checked_not_just_style_config(self):
        from copy import deepcopy
        from docx.oxml.ns import qn
        path=self.project/"reference-numbering.docx";doc=Document();doc.add_heading("Reference style validation",0);doc.add_heading("References",1)
        paragraph=doc.add_paragraph("Internal validation reference.")
        paragraph._p.get_or_add_pPr().append(deepcopy(doc.styles["List Number"].element.find(qn("w:pPr")).find(qn("w:numPr"))))
        doc.save(path);normalize_docx(path,DEFAULT_STYLE,"title_page")
        correct=audit_docx(path,{**DEFAULT_STYLE,"reference_number_separator":"."},"manuscript")
        wrong=audit_docx(path,{**DEFAULT_STYLE,"reference_number_separator":")"},"manuscript")
        self.assertFalse(any("separator" in message for message in correct),correct)
        self.assertTrue(any("separator differs" in message for message in wrong),wrong)

    def test_external_builder_input_drift_invalidates_round_proof(self):
        source,_,output=self.baseline()
        with tempfile.TemporaryDirectory(prefix="medpaper-external-style-") as temp:
            external=Path(temp)/"style.csl";external.write_text("Original citation style")
            revision.record_build(self.project,output,[source,external],entrypoints=[source],tool="Actual test producer")
            _,row=self.round([source.relative_to(self.project).as_posix()])
            before=revision.input_signature(self.project,row)
            external.write_text("Changed citation style")
            self.assertNotEqual(before,revision.input_signature(self.project,row))


if __name__=="__main__":unittest.main()
